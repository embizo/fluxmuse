"""FluxMuse financial model v2: workbook writer (openpyxl)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

import fm_annual as A
import fm_inputs as I
import fm_rows as F

VERSION = "v2"
MODEL_DATE = "2026-09-11"
ORANGE, SLATE, NIGHT, TINT, MIST = "FF6A00", "37474F", "0F1419", "FFF4EB", "F3F4F6"
FONT = "Arial"
FMT = {
    "zar": '"R"#,##0;("R"#,##0);"-"', "usd": '"US$"#,##0;("US$"#,##0);"-"',
    "n0": '#,##0;(#,##0);"-"', "n1": '#,##0.0;(#,##0.0);"-"', "n2": '#,##0.00;(#,##0.00);"-"',
    "pct": '0.0%;(0.0%);"-"', "pct2": '0.00%;(0.00%);"-"', "x1": '0.0"x";(0.0"x");"-"', "x2": '0.00"x"',
    "x3": '0.000"x"', "int": '0', "num": '#,##0.##', "num1": '#,##0.0', "num2": '#,##0.00', "text": '@',
    "ngn": '"₦"#,##0', "kes": '"KSh "#,##0', "ghs": '"GH₵ "#,##0',
}
F_TITLE = Font(name=FONT, size=16, bold=True, color=NIGHT)
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10, color="000000")
F_BOLD = Font(name=FONT, size=10, bold=True, color="000000")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_BUILD = Font(name=FONT, size=10, italic=True, color="595959")
F_FLAG = Font(name=FONT, size=10, bold=True, color="C62828")
FILL_HDR = PatternFill("solid", fgColor=ORANGE)
FILL_SUB = PatternFill("solid", fgColor=SLATE)
FILL_INPUT = PatternFill("solid", fgColor="FFF9DB")
FILL_BUILD = PatternFill("solid", fgColor=MIST)
FILL_TOTAL = PatternFill("solid", fgColor=TINT)
WRAP = Alignment(wrap_text=True, vertical="top")


def _cell(ws, addr, value, font=F_BODY, fmt=None, fill=None, align=None):
    c = ws[addr]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def _title(ws, title, subtitle):
    _cell(ws, "A1", title, F_TITLE)
    _cell(ws, "A2", subtitle, F_SUB)


def _hdr_row(ws, row, c1, c2, fill=FILL_HDR, font=F_HDR):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font


def _input(ws, addr, value, fmt):
    return _cell(ws, addr, value, F_INPUT, fmt, FILL_INPUT)


def _src_font(src):
    return F_FLAG if ("[[" in str(src) or "NOT IN CURRENT PRICING" in str(src)) else F_SUB


# ---------------------------------------------------------------------------
# layout (row numbers must exist before any formula text is generated)
# ---------------------------------------------------------------------------
MONTHLY_SHEETS = ["Revenue_Build", "Costs", "P&L", "Cash_Flow"]
ANNUAL_SHEETS = ["Annual_Summary", "Unit_Economics"]
SCALAR_START = {}


def layout():
    nxt = {s: F.FIRST_ROW for s in MONTHLY_SHEETS}
    for rw in F.ROWS:
        s = rw["sheet"]
        if rw["kind"] == "section":
            if nxt[s] > F.FIRST_ROW:
                nxt[s] += 1
            rw["row"] = nxt[s]
        else:
            rw["row"] = nxt[s]
            F.ROWPOS[rw["key"]] = (s, nxt[s])
        nxt[s] += 1
    SCALAR_START["Cash_Flow"] = nxt["Cash_Flow"] + 2
    r = SCALAR_START["Cash_Flow"] + 1
    for s in A.SROWS:
        s["row"] = r
        A.SCALPOS[s["key"]] = r
        r += 1
    anx = {s: A.ANN_FIRST_ROW for s in ANNUAL_SHEETS}
    for rw in A.AROWS:
        s = rw["sheet"]
        if rw["kind"] == "section":
            if anx[s] > A.ANN_FIRST_ROW:
                anx[s] += 1
            rw["row"] = anx[s]
        else:
            rw["row"] = anx[s]
            A.ANNPOS[rw["key"]] = (s, anx[s])
        anx[s] += 1


# ---------------------------------------------------------------------------
# Assumptions & Pricing
# ---------------------------------------------------------------------------
def write_assumptions(wb):
    ws = wb.create_sheet("Assumptions")
    _title(ws, "Assumptions", "Every driver of the model. Blue on yellow = input. Black = formula. Red source text = [[CONFIRM]] / founder decision pending.")
    heads = ["Driver", "Live value", "Unit", "Conservative", "Base", "Upside", "Source / rationale"]
    for k, h in enumerate(heads):
        ws.cell(row=4, column=k + 1, value=h)
    _hdr_row(ws, 4, 1, 7)
    _cell(ws, "A6", "Active scenario (select)", F_BOLD)
    _input(ws, "B6", "Base", "text")
    _cell(ws, "C6", "list")
    _cell(ws, "G6", "Choose Conservative / Base / Upside. Every scenario driver below uses CHOOSE on this index.", F_SUB)
    dv = DataValidation(type="list", formula1='"Conservative,Base,Upside"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B6")
    _cell(ws, "A7", "Scenario index (1 = Conservative, 2 = Base, 3 = Upside)", F_BODY)
    _cell(ws, "B7", "=MATCH($B$6,$D$4:$F$4,0)", F_BODY, "int")
    for e in I.ENTRIES:
        r = e["row"]
        if e["kind"] == "section":
            ws.cell(row=r, column=1, value=e["title"])
            _hdr_row(ws, r, 1, 7, FILL_SUB)
            continue
        if e["kind"] == "note":
            _cell(ws, f"A{r}", e["title"], F_SUB)
            continue
        _cell(ws, f"A{r}", e["label"])
        _cell(ws, f"C{r}", e["unit"])
        _cell(ws, f"G{r}", e["src"], _src_font(e["src"]))
        if e["kind"] == "inp":
            _input(ws, f"B{r}", e["value"], e["fmt"])
        elif e["kind"] == "scen":
            for k, colL in enumerate("DEF"):
                _input(ws, f"{colL}{r}", e["values"][k], e["fmt"])
            _cell(ws, f"B{r}", f"=CHOOSE($B$7,D{r},E{r},F{r})", F_BODY, e["fmt"])
        else:
            _cell(ws, f"B{r}", "=" + e["xl"](I.AD), F_BODY, e["fmt"])
    hr = I.HC_HEADER_ROW
    ws.cell(row=hr - 1, column=1, value="Headcount plan: monthly cost-to-company in 2026 rand (escalated). COST DISCIPLINE: a role is hired only once month # >= earliest hire, the seed has landed, and last month's net MRR >= MRR gate x scenario gate multiplier")
    _hdr_row(ws, hr - 1, 1, 7, FILL_SUB)
    for k, h in enumerate(["Role", "Monthly CTC (R)", "Department", "FTE count", "Earliest hire month #", "MRR gate (R net MRR)", "Note"]):
        ws.cell(row=hr, column=k + 1, value=h)
    _hdr_row(ws, hr, 1, 7)
    for i, ro in enumerate(I.ROLES):
        r = ro["row"]
        _cell(ws, f"A{r}", ro["name"])
        _input(ws, f"B{r}", ro["ctc"], "zar")
        _cell(ws, f"C{r}", ro["dept"])
        _input(ws, f"D{r}", ro["count"], "int")
        _input(ws, f"E{r}", ro["hire"], "int")
        _input(ws, f"F{r}", ro["gate"], "zar")
        if ro["type"] == "F":
            note = "Founder: from month 1 (lean pre-seed mode)"
        elif ro["link"]:
            mk, kind = ro["link"]
            note = {"lead": f"Hired at the {I.MARKET_SHORT[mk]} launch decision (2 months before launch)",
                    "launch": f"Hired in the {I.MARKET_SHORT[mk]} launch month",
                    "exp": f"Hired 12 months after the {I.MARKET_SHORT[mk]} launch"}[kind] + " (gate & earliest-hire columns not used)"
        else:
            note = f"Earliest {F.mlabel(ro['hire'])}" + ("; partnerships manager (drives agency sign-ups)" if ro["type"] == "PM" else "")
        _cell(ws, f"G{r}", note, F_SUB)
    ws.column_dimensions["A"].width = 64
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 24
    for colL in "DEF":
        ws.column_dimensions[colL].width = 15
    ws.column_dimensions["G"].width = 95
    ws.freeze_panes = "B5"


def write_pricing(wb):
    ws = wb.create_sheet("Pricing")
    _title(ws, "Pricing", "Live checkout prices (facts file §2). ZAR list for South Africa; fixed local price points for NG/KE/GH and USD for the rest of rail-covered Africa, at FX parity. Annual = 10x monthly.")
    heads = ["Tier", "ZAR list / month", "ZAR annual", "≈ US$ / month (R18.50)", "Included AI credits / month",
             "Founding Member price, first 2 monthly bills (R)", "Pilot price, first 2 monthly bills (R)",
             "NGN / month", "KES / month", "GHS / month", "USD / month (19 USD markets)", "Note"]
    for k, h in enumerate(heads):
        ws.cell(row=4, column=k + 1, value=h)
    _hdr_row(ws, 4, 1, 12)
    ws.row_dimensions[4].height = 44
    for k in range(1, 13):
        ws.cell(row=4, column=k).alignment = Alignment(wrap_text=True, vertical="center")
    lfmt = {"H": "ngn", "I": "kes", "J": "ghs", "K": "usd"}
    for i, t in enumerate(I.PRICE_KEYS):
        r = I.PRICING_TIER_ROW0 + i
        if t == "W":
            _cell(ws, f"A{r}", "Partner wholesale: Agency tier -30%", F_BOLD)
            _input(ws, f"B{r}", I.WHOLESALE[0], "zar")
            locals_ = I.WHOLESALE[1]
            _cell(ws, f"E{r}", f"=E{I.PRICING_TIER_ROW0 + 3}", F_BODY, "n0")
            _cell(ws, f"F{r}", "Does not stack", F_SUB)
            note = "Founder decision 2026-09-11: partners/agencies buy Agency at 30% off list and set their own retail price. Referral commission: undecided, not modelled."
        else:
            price, credits, note = I.PRICING[t]
            _cell(ws, f"A{r}", I.TIER_NAME[t], F_BOLD)
            _input(ws, f"B{r}", price, "zar")
            _input(ws, f"E{r}", credits, "n0")
            locals_ = I.LOCAL[t]
            if t in I.FM_PRICE:
                _cell(ws, f"F{r}", f"=ROUNDDOWN(B{r}*(1-{I.AD('fmB_disc')}),0)", F_BODY, "zar")
            else:
                _cell(ws, f"F{r}", "Not offered" if t == "A" else "n/a", F_SUB)
            if t in I.PILOT_PRICE:
                _cell(ws, f"G{r}", f"=ROUNDDOWN(B{r}*(1-{I.AD('pilot_disc')}),0)", F_BODY, "zar")
            else:
                _cell(ws, f"G{r}", "n/a", F_SUB)
        _cell(ws, f"C{r}", f"=B{r}*{I.AD('annual_months')}", F_BODY, "zar")
        _cell(ws, f"D{r}", f"=B{r}/{I.AD('fx')}", F_BODY, "usd")
        for colL, val in zip("HIJK", locals_):
            _input(ws, f"{colL}{r}", val, lfmt[colL])
        _cell(ws, f"L{r}", note, F_SUB)
    r0 = I.PRICING_SCALAR_ROW0
    ws.cell(row=r0 - 1, column=1, value="Pricing drivers")
    _hdr_row(ws, r0 - 1, 1, 12, FILL_SUB)
    for i, (key, label, unit, val, note, fmt) in enumerate(I.PRICING_SCALARS):
        r = r0 + i
        _cell(ws, f"A{r}", label)
        if val is None:
            _cell(ws, f"B{r}", f"=1-{I.AD('annual_share')}*(12-{I.AD('annual_months')})/12", F_BODY, fmt)
        else:
            _input(ws, f"B{r}", val, fmt)
        _cell(ws, f"C{r}", unit)
        _cell(ws, f"D{r}", note, F_SUB)
    rf = I.FX_ROW0
    ws.cell(row=rf - 1, column=1, value="FX: price-setting rates, depreciation vs ZAR and the quarterly repricing policy (depreciation rates are ASSUMPTIONS)")
    _hdr_row(ws, rf - 1, 1, 12, FILL_SUB)
    for i, (key, label, unit, val, note, fmt) in enumerate(I.FX_ROWS):
        r = rf + i
        _cell(ws, f"A{r}", label)
        _input(ws, f"B{r}", val, fmt)
        _cell(ws, f"C{r}", unit)
        _cell(ws, f"D{r}", note, F_FLAG if key.startswith("dep_") else F_SUB)
    rp = I.PAR_HDR_ROW
    ws.cell(row=rp, column=1, value="ZAR value of each market's monthly price at price-setting FX (formulas; used by Revenue_Build, then x FX value factor after drift & repricing)")
    _hdr_row(ws, rp, 1, 12, FILL_SUB)
    for k, h in enumerate(["Tier (R / month)"] + [I.MARKET_SHORT[r] for r in I.PAR_COLS]):
        ws.cell(row=rp + 1, column=k + 1, value=h)
    _hdr_row(ws, rp + 1, 1, 7)
    for i, t in enumerate(I.PRICE_KEYS):
        r = I.PAR_ROW0 + i
        _cell(ws, f"A{r}", "Partner wholesale (Agency -30%)" if t == "W" else I.TIER_NAME[t], F_BOLD)
        for mk, colL in I.PAR_COLS.items():
            _cell(ws, f"{colL}{r}", "=" + I.par_price_xl(mk, t), F_BODY, "zar")
    ws.column_dimensions["A"].width = 58
    for colL in "BCDEFGHIJK":
        ws.column_dimensions[colL].width = 16
    ws.column_dimensions["L"].width = 80
    ws.freeze_panes = "B5"


# ---------------------------------------------------------------------------
# monthly sheets
# ---------------------------------------------------------------------------
SUBTITLES = {
    "Revenue_Build": "Monthly, Oct 2026 - Sep 2031. Launch gates, FX & repricing, acquisition by market & segment with the marketing cap, cohorts, gross -> discounts -> net revenue.",
    "Costs": "Monthly MRR-gated headcount plan, COGS and operating expenses (lean pre-seed mode until the seed lands).",
    "P&L": "Monthly profit and loss (ZAR). EBITDA excludes D&A (immaterial: equipment is expensed).",
    "Cash_Flow": "Monthly cash flow & funding (R25M seed only), R3.0M buffer test, runway and key outputs.",
}


def write_monthly(wb):
    sheets = {s: wb.create_sheet(s) for s in MONTHLY_SHEETS}
    for s, ws in sheets.items():
        _title(ws, s.replace("_", " "), SUBTITLES[s])
        _cell(ws, "A3", "Month #", F_HDR, None, FILL_HDR)
        _cell(ws, "A4", "Month", F_HDR, None, FILL_HDR)
        _cell(ws, "B3", "Unit", F_HDR, None, FILL_HDR)
        _cell(ws, "B4", "", F_HDR, None, FILL_HDR)
        _cell(ws, "A5", "Fiscal year (Oct-Sep)", F_BOLD)
        for m in F.MONTHS:
            c = F.col(m)
            _cell(ws, f"{c}3", m, F_HDR, "int", FILL_HDR, Alignment(horizontal="center"))
            _cell(ws, f"{c}4", F.mlabel(m), F_HDR, "text", FILL_HDR, Alignment(horizontal="center"))
            _cell(ws, f"{c}5", F.fy_of(m), F_BODY, "int", None, Alignment(horizontal="center"))
        ws.column_dimensions["A"].width = 70
        ws.column_dimensions["B"].width = 12
        for m in F.MONTHS:
            ws.column_dimensions[F.col(m)].width = 12.5
        ws.freeze_panes = "C6"
    for rw in F.ROWS:
        s = rw["sheet"]
        ws = sheets[s]
        F.CUR[0] = s
        r = rw["row"]
        if rw["kind"] == "section":
            ws.cell(row=r, column=1, value=rw["label"])
            _hdr_row(ws, r, 1, 2 + F.M, FILL_SUB)
            continue
        font = F_BOLD if rw["bold"] else F_BODY
        _cell(ws, f"A{r}", rw["label"], font)
        _cell(ws, f"B{r}", rw["unit"], F_SUB)
        for m in F.MONTHS:
            c = _cell(ws, f"{F.col(m)}{r}", "=" + rw["xl"](m), font, rw["fmt"])
            if rw["bold"]:
                c.fill = FILL_TOTAL
    ws = sheets["Cash_Flow"]
    r0 = SCALAR_START["Cash_Flow"]
    ws.cell(row=r0, column=1, value="KEY CASH OUTPUTS (selected scenario)")
    _hdr_row(ws, r0, 1, 2)
    F.CUR[0] = "Cash_Flow"
    for s in A.SROWS:
        _cell(ws, f"A{s['row']}", s["label"], F_BOLD if not s["key"].endswith("_num") else F_BODY)
        _cell(ws, f"B{s['row']}", "=" + s["xl"](), F_BOLD, s["fmt"], FILL_TOTAL, Alignment(horizontal="right"))
    ws.column_dimensions["B"].width = 30
    F.CUR[0] = None


# ---------------------------------------------------------------------------
# annual sheets
# ---------------------------------------------------------------------------
def write_annual(wb, sheet_name, subtitle):
    ws = wb.create_sheet(sheet_name)
    _title(ws, sheet_name.replace("_", " "), subtitle)
    for y in A.YEARS:
        c = A.acol(y)
        _cell(ws, f"{c}4", f"FY{y}", F_HDR, "text", FILL_HDR, Alignment(horizontal="center"))
        _cell(ws, f"{c}5", f"Oct {2025 + y} - Sep {2026 + y}", F_SUB, "text", None, Alignment(horizontal="center"))
    _cell(ws, "A4", "R unless stated", F_HDR, None, FILL_HDR)
    _cell(ws, "B4", "Unit", F_HDR, None, FILL_HDR)
    A.ACUR[0] = sheet_name
    for rw in A.AROWS:
        if rw["sheet"] != sheet_name:
            continue
        r = rw["row"]
        if rw["kind"] == "section":
            ws.cell(row=r, column=1, value=rw["label"])
            _hdr_row(ws, r, 1, 7, FILL_SUB)
            continue
        font = F_BOLD if rw["bold"] else F_BODY
        _cell(ws, f"A{r}", rw["label"], font)
        _cell(ws, f"B{r}", rw["unit"], F_SUB)
        for y in A.YEARS:
            c = _cell(ws, f"{A.acol(y)}{r}", "=" + rw["xl"](y), font, rw["fmt"], None, Alignment(horizontal="right"))
            if rw["bold"]:
                c.fill = FILL_TOTAL
    A.ACUR[0] = None
    ws.column_dimensions["A"].width = 66
    ws.column_dimensions["B"].width = 22
    for y in A.YEARS:
        ws.column_dimensions[A.acol(y)].width = 17
    ws.freeze_panes = "C6"
    return ws


def annual_ref(key, y):
    sh, r = A.ANNPOS[key]
    return f"{F.sref(sh)}!${A.acol(y)}${r}"


def month_sum(key, m1, m2):
    sh, r = F.ROWPOS[key]
    return f"SUM({F.sref(sh)}!{F.col(m1)}{r}:{F.col(m2)}{r})"


# ---------------------------------------------------------------------------
# Use of funds
# ---------------------------------------------------------------------------
USE_OF_FUNDS = [
    ("Product & engineering", 0.40,
     "Senior full-stack x2, AI/ML, designer, full-stack x2, DevOps/QA, integrations x2, PM, security; founders' salaries; "
     "hosting & AI inference; Instagram/Threads/LinkedIn/TikTok adapters, commerce/CRM integrations, AI cost optimisation."),
    ("Sales & marketing / partner programme", 0.30,
     "Head of Growth, performance & multilingual content marketers, partnerships managers, SME onboarding; paid acquisition "
     "(capped at a % of MRR and a CAC-payback test), brand & community, Flux_Partner programme."),
    ("Market expansion (NG, KE, GH) & payments compliance", 0.15,
     "Country leads and local sales/CS in Nigeria, Kenya and Ghana; NDPR/KDPA/Ghana DPC registrations, local counsel, "
     "launch campaigns, POPIA/Meta compliance base, Rest-of-Africa VAT registrations."),
    ("Operations & working capital", 0.15,
     "Customer success team, finance & compliance, operations, G&A, tooling, travel, payment processing, support, WhatsApp "
     "pass-through costs, and the working-capital float that protects the R3.0M cash buffer."),
]
# (category index) -> (plus keys, minus keys)
UOF_MAP = [
    (["cost_dept_Engineering & product", "cost_dept_Leadership", "cogs_hosting", "cogs_ai", "opex_tools"], []),
    (["pl_mkt", "cost_dept_Marketing", "cost_dept_Sales & partnerships"], ["opex_launch_mkt", "opex_rest_mkt"]),
    (["cost_dept_Country teams", "opex_legal", "opex_launch_mkt", "opex_rest_mkt"], []),
    (["cost_dept_Customer success", "cost_dept_G&A & compliance", "opex_ga", "opex_travel", "cogs_support", "cogs_proc",
      "cogs_wa", "tax", "d_ar"], ["d_dr"]),
]
UOF_WINDOWS = (18, 24)


def uof_window(p, n):
    s = int(p["seed_month"])
    return s, min(F.M, s + n - 1)


def uof_xl(i, m1, m2):
    plus_, minus_ = UOF_MAP[i]
    return "+".join(month_sum(k, m1, m2) for k in plus_) + "".join("-" + month_sum(k, m1, m2) for k in minus_)


def uof_py(i, m1, m2):
    plus_, minus_ = UOF_MAP[i]
    return sum(sum(F.V[k][m1:m2 + 1]) for k in plus_) - sum(sum(F.V[k][m1:m2 + 1]) for k in minus_)


UOF_CELLS = {}   # (window n, category index or 'total'/'rev'/'net') -> address, for verification


def write_use_of_funds(wb, p):
    ws = wb.create_sheet("Use_of_Funds")
    _title(ws, "Use of Funds", "Seed R25M split (facts file §7). Percentages are inputs; amounts are formulas on the seed amount.")
    for k, h in enumerate(["Category", "Share", "Amount (R)", "≈ US$", "What it buys (hires & milestones)"]):
        ws.cell(row=4, column=k + 1, value=h)
    _hdr_row(ws, 4, 1, 5)
    for i, (cat, pct, buys) in enumerate(USE_OF_FUNDS):
        r = 5 + i
        _cell(ws, f"A{r}", cat, F_BOLD)
        _input(ws, f"B{r}", pct, "pct")
        _cell(ws, f"C{r}", f"={I.AD('seed_amount')}*B{r}", F_BODY, "zar")
        _cell(ws, f"D{r}", f"=C{r}/{I.AD('fx')}", F_BODY, "usd")
        _cell(ws, f"E{r}", buys, F_BODY, None, None, WRAP)
        ws.row_dimensions[r].height = 58
    _cell(ws, "A9", "Total", F_BOLD, None, FILL_TOTAL)
    _cell(ws, "B9", "=SUM(B5:B8)", F_BOLD, "pct", FILL_TOTAL)
    _cell(ws, "C9", "=SUM(C5:C8)", F_BOLD, "zar", FILL_TOTAL)
    _cell(ws, "D9", "=SUM(D5:D8)", F_BOLD, "usd", FILL_TOTAL)
    _cell(ws, "E9", '=IF(ABS(B9-1)<0.0001,"Check: shares sum to 100%","CHECK: shares do not sum to 100%")', F_BOLD, None, FILL_TOTAL)
    r = 11
    for n in UOF_WINDOWS:
        m1, m2 = uof_window(p, n)
        ws.cell(row=r, column=1, value=f"Reconciliation: modelled gross spend in the {n} months after close (months {m1}-{m2}: {F.mlabel(m1)} - {F.mlabel(m2)}; window fixed at build time from the seed month)")
        _hdr_row(ws, r, 1, 5, FILL_SUB)
        r += 1
        for k, h in enumerate(["Category", "Allocation share", "Modelled spend (R)", "Share of modelled spend", "Gap vs allocation (share of spend - allocation, pp)"]):
            ws.cell(row=r, column=k + 1, value=h)
        _hdr_row(ws, r, 1, 5)
        r += 1
        first = r
        for i, (cat, pct, buys) in enumerate(USE_OF_FUNDS):
            _cell(ws, f"A{r}", cat)
            _cell(ws, f"B{r}", f"=B{5 + i}", F_BODY, "pct")
            _cell(ws, f"C{r}", "=" + uof_xl(i, m1, m2), F_BODY, "zar")
            _cell(ws, f"D{r}", f"=IF(SUM($C${first}:$C${first + 3})=0,0,C{r}/SUM($C${first}:$C${first + 3}))", F_BODY, "pct")
            _cell(ws, f"E{r}", f"=D{r}-B{r}", F_BODY, "pct")
            UOF_CELLS[(n, i)] = f"C{r}"
            r += 1
        _cell(ws, f"A{r}", "Total gross spend (COGS + opex + tax + working capital)", F_BOLD, None, FILL_TOTAL)
        _cell(ws, f"C{r}", f"=SUM(C{first}:C{first + 3})", F_BOLD, "zar", FILL_TOTAL)
        UOF_CELLS[(n, "total")] = f"C{r}"
        tot = r
        r += 1
        _cell(ws, f"A{r}", "less: revenue collected in the window")
        _cell(ws, f"C{r}", "=" + month_sum("pl_rev", m1, m2), F_BODY, "zar")
        UOF_CELLS[(n, "rev")] = f"C{r}"
        r += 1
        _cell(ws, f"A{r}", "Net cash consumed (funded by the seed)", F_BOLD)
        _cell(ws, f"C{r}", f"=C{tot}-C{tot + 1}", F_BOLD, "zar")
        UOF_CELLS[(n, "net")] = f"C{r}"
        r += 1
        _cell(ws, f"A{r}", "Seed remaining unspent at the end of the window (seed - net cash consumed)")
        _cell(ws, f"C{r}", f"={I.AD('seed_amount')}-C{tot + 2}", F_BODY, "zar")
        r += 2
    _cell(ws, f"A{r}", "Reading: the seed funds net burn, not gross spend, so category shares of gross spend are compared with the allocation shares. "
          "Revenue funds the balance. A positive gap means the category takes more of the modelled spend than its allocation.", F_SUB)
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 80
    ws.freeze_panes = "A5"


def write_market_sizing(wb):
    ws = wb.create_sheet("Market_Sizing")
    _title(ws, "Market Sizing", "TAM / SAM / SOM from facts file §5. All figures are estimates (ASSUMPTION): cite and verify before external use.")
    for k, h in enumerate(["Measure", "Value", "Unit", "Source / basis"]):
        ws.cell(row=4, column=k + 1, value=h)
    _hdr_row(ws, 4, 1, 4)
    rows = [
        ("TAM: MSMEs in Sub-Saharan Africa", 44000000, "businesses", "IFC / World Bank est. (~44M MSMEs)", True, "n0"),
        ("SAM: digitally active SMBs in the 23 rail-covered countries selling via social/WhatsApp, able to pay >= US$25/mo", 3500000, "businesses", "Planning estimate (facts file §5)", True, "n0"),
        ("SOM share of SAM (5-year)", 0.005, "%", "Planning estimate: ~0.5% of SAM", True, "pct"),
        ("SOM: paying workspaces by FY5", "=B6*B7", "workspaces", "Formula: SAM x SOM share", False, "n0"),
        ("Minimum price point used for SAM value", 25, "US$ / month", "SAM definition: can pay >= US$25/mo", True, "usd"),
        ("TAM value at US$25/mo (illustrative ceiling)", f"=B5*B9*12*{I.AD('fx')}", "R / year", "Overstates: most informal MSMEs will not pay", False, "zar"),
        ("SAM value at US$25/mo", f"=B6*B9*12*{I.AD('fx')}", "R / year", "", False, "zar"),
        ("SOM value at modelled FY5 blended subscription ARPA", f"=B8*{annual_ref('a_arpa', 5)}*12", "R / year", "Links to Annual_Summary FY5 ARPA", False, "zar"),
        ("Model: paying workspaces at end of FY5 (selected scenario)", f"={annual_ref('a_cust', 5)}", "workspaces", "Links to Annual_Summary", False, "n0"),
        ("Model FY5 workspaces as % of SOM", "=IF(B8=0,0,B13/B8)", "%", "", False, "pct"),
        ("Model FY5 workspaces as % of SAM", "=IF(B6=0,0,B13/B6)", "%", "", False, "pct2"),
    ]
    for i, (lab, val, unit, src, is_in, fmt) in enumerate(rows):
        r = 5 + i
        _cell(ws, f"A{r}", lab, F_BOLD if not is_in else F_BODY)
        if is_in:
            _input(ws, f"B{r}", val, fmt)
        else:
            _cell(ws, f"B{r}", val, F_BODY, fmt)
        _cell(ws, f"C{r}", unit)
        _cell(ws, f"D{r}", src, F_SUB)
    ws.cell(row=18, column=1, value="Country context (estimates, for narrative only). Gated countries (no rail) are zero in every scenario.")
    _hdr_row(ws, 18, 1, 4, FILL_SUB)
    ctx = [("South Africa SMMEs", "2.5-3 million", "SEDA / Stats SA"), ("Nigeria MSMEs", "~39 million", "SMEDAN"),
           ("Kenya MSMEs", "~7.4 million", "KNBS"),
           ("WhatsApp reach", ">90% of internet users in SA and NG", "DataReportal 2025"),
           ("Mobile money", "SSA ~70% of global mobile money value; >US$1 trillion processed / year", "GSMA State of the Industry 2025")]
    for i, (a, b, c) in enumerate(ctx):
        r = 19 + i
        _cell(ws, f"A{r}", a)
        _cell(ws, f"B{r}", b)
        _cell(ws, f"D{r}", c, F_SUB)
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 50


# ---------------------------------------------------------------------------
# Sensitivity (build-time values)
# ---------------------------------------------------------------------------
def write_sensitivity(wb, sens):
    ws = wb.create_sheet("Sensitivity")
    _title(ws, "Sensitivity", "VALUES COMPUTED AT BUILD TIME by the Python mirror of this workbook (grey). They do not recalculate; rerun _build/financial_model.py.")
    r = 4
    for blk in sens["blocks"]:
        ncol = len(blk["cols"]) + 1
        ws.cell(row=r, column=1, value=blk["title"])
        _hdr_row(ws, r, 1, max(ncol, 2))
        r += 1
        for k, h in enumerate(["Output"] + blk["cols"]):
            c = ws.cell(row=r, column=k + 1, value=h)
            c.alignment = Alignment(wrap_text=True, horizontal="center" if k else "left")
        _hdr_row(ws, r, 1, ncol, FILL_SUB)
        ws.row_dimensions[r].height = 30
        for lab, fmt, vals in blk["rows"]:
            r += 1
            _cell(ws, f"A{r}", lab)
            for k, val in enumerate(vals):
                colL = F.get_column_letter(2 + k)
                _cell(ws, f"{colL}{r}", val, F_BUILD, fmt if not isinstance(val, str) else "text", FILL_BUILD, Alignment(horizontal="right"))
        if blk.get("note"):
            r += 1
            _cell(ws, f"A{r}", blk["note"], F_SUB)
        r += 2
    g = sens["grid"]
    for tbl in g["tables"]:
        ws.cell(row=r, column=1, value=f"Base: {tbl['title']} (computed at build time)")
        _hdr_row(ws, r, 1, 7)
        r += 1
        _cell(ws, f"A{r}", "Churn multiplier ↓  /  New-customer volume multiplier →", F_BOLD)
        for j, vm in enumerate(g["vol"]):
            _cell(ws, f"{'CDEFG'[j]}{r}", vm, F_HDR, "x2", FILL_SUB, Alignment(horizontal="center"))
        for i, cm in enumerate(g["churn"]):
            r += 1
            _cell(ws, f"A{r}", f"{cm:.2f}x churn (Solo Starter {cm * g['base_starter_churn']:.1%}/mo)")
            _cell(ws, f"B{r}", cm, F_BUILD, "x2")
            for j in range(len(g["vol"])):
                c = _cell(ws, f"{'CDEFG'[j]}{r}", tbl["values"][i][j], F_BUILD, tbl["fmt"], FILL_BUILD)
                if cm == 1.0 and g["vol"][j] == 1.0:
                    c.font = Font(name=FONT, size=10, bold=True, color=NIGHT)
                    c.fill = FILL_TOTAL
        r += 2
    ws.column_dimensions["A"].width = 66
    for colL in "BCDEFGH":
        ws.column_dimensions[colL].width = 19


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------
CHANGES_FROM_V1 = [
    ("Seed & timing", "Seed R18.5M (Dec 2026) -> R25M landing Feb 2027 (input). Lean pre-seed mode Oct 2026-Jan 2027: founders, infrastructure and pilot costs only. New output: pre-seed bridge required."),
    ("Profitability constraint", "R25M must reach profitability alone: no Series A in Base or Conservative (input kept at 0). Closing cash >= R3.0M buffer every month from the seed month; sustained EBITDA and operating-cash-flow break-even."),
    ("Cost discipline", "Hires and NG/KE/GH/Rest launches gated on last month's net MRR x a scenario gate multiplier (Conservative 1.40, Base 1.00, Upside 0.80). Paid acquisition = MIN(desired, floor + % of MRR) and switched off where CAC payback > 12 months. Brand budget scales with MRR. Hiring now differs by scenario."),
    ("Pilot", "12 Gauteng brands free in Oct-Nov 2026; 75% convert on 1 Dec 2026 at 50% off the first 2 bills (R249/R999/R2,499/R3,999). No other paying customers before the 1 Dec SA launch."),
    ("Founding Member", "30% off the first 2 monthly bills for sign-ups in each launch window (ZA Dec 2026-Jan 2027; NG/KE/GH first 60 days); annual plans get 2 bonus months (fee recognised over 14 months). Not on Agency. +20% sign-up uplift in windows. Option A / None as sensitivities."),
    ("Markets", "v1 regions removed. ZA -> Nigeria, Kenya, Ghana (own rows, launch month, country lead, launch marketing, compliance) -> Rest of rail-covered Africa (19 countries, USD, self-serve, no team) -> Botswana & Namibia (off). Gated countries = 0."),
    ("Pricing", "v1's 75-85% regional price index removed. Local NGN/KES/GHS and USD price points at FX parity, with depreciation vs ZAR and quarterly repricing when drift > 10%. FX-shock sensitivity (ZAR 15% stronger)."),
    ("Segments", "v1 channels (self-serve / agency / direct AE) replaced by segments: Solo (Starter -> Growth), SMEs (Growth -> Scale), Agencies, Enterprise inbound only. Own CAC, churn, mix and upgrade rates. Account executives removed."),
    ("Partner wholesale", "Agency-segment customers pay partner wholesale R5,599/mo (₦463,000 / KSh 44,999 / GH₵ 3,799 / $299). v1's 40% sub-account wholesale discount and separate client sub-account revenue removed. Referral commission not modelled (undecided)."),
    ("Commerce fee", "0.75% commerce fee now Upside only (0% Base & Conservative); flagged not in current pricing. Campaign Financing still excluded."),
    ("Use of funds", "R25M: 40/30/15/15, reconciled against modelled spend in the 18 and 24 months after close."),
]


def write_cover(wb):
    ws = wb["Cover"]
    _cell(ws, "A1", "FluxMuse Financial Model", Font(name=FONT, size=22, bold=True, color=ORANGE))
    _cell(ws, "A2", "AI marketing team & WhatsApp commerce for African SMBs | Fluxmuse Pty Ltd", Font(name=FONT, size=11, color=SLATE))
    info = [("Version", f"{VERSION} (R25M seed, profitable on the seed alone)"), ("Model date", MODEL_DATE),
            ("Horizon", "60 months: Oct 2026 - Sep 2031 (FY1-FY5, fiscal year Oct-Sep). Month 1 = Oct 2026"),
            ("Currency", "ZAR (R); US$ at R18.50 = US$1 (input on Assumptions)")]
    for k, (a, b) in enumerate(info):
        _cell(ws, f"A{4 + k}", a, F_BOLD)
        _cell(ws, f"B{4 + k}", b)
    _cell(ws, "A8", "Active scenario", F_BOLD)
    _cell(ws, "B8", "=Assumptions!$B$6", F_BOLD)
    _cell(ws, "A9", "Disclaimer", F_BOLD)
    _cell(ws, "B9", "Forward-looking projections; assumptions to be validated with pilot data. Red-flagged inputs need founder confirmation.", F_FLAG)
    ws.cell(row=11, column=1, value="Headline (live formulas, selected scenario)")
    _hdr_row(ws, 11, 1, 7)
    for y in A.YEARS:
        _cell(ws, f"{'CDEFG'[y - 1]}11", f"FY{y}", F_HDR, None, FILL_HDR, Alignment(horizontal="center"))
    heads = [("Total revenue (net of launch discounts)", "a_rev", "zar"), ("EBITDA", "a_ebitda", "zar"), ("EBITDA margin", "a_ebitda_pct", "pct"),
             ("Paying workspaces (end of FY)", "a_cust", "n0"), ("Subscription ARR", "a_arr", "zar"),
             ("Closing cash", "a_close_cash", "zar"), ("Headcount (end of FY)", "a_fte", "n0"),
             ("Founding Member discount cost", "a_fm_cost", "zar")]
    r = 12
    for lab, key, fmt in heads:
        _cell(ws, f"A{r}", lab)
        for y in A.YEARS:
            _cell(ws, f"{'CDEFG'[y - 1]}{r}", f"={annual_ref(key, y)}", F_BODY, fmt)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Constraint check & key outputs (live formulas, selected scenario)")
    _hdr_row(ws, r, 1, 7)
    for lab, key, fmt in [("Profitable on the R25M seed alone?", "k_profitable", None),
                          ("EBITDA break-even (first month)", "k_be", None), ("EBITDA break-even (sustained)", "k_be_sus", None),
                          ("Operating cash flow positive (first month)", "k_cfpos", None), ("Operating cash flow positive (sustained)", "k_cfpos_sus", None),
                          ("Minimum closing cash from the seed month on", "k_min_post", "zar"), ("Month of minimum post-seed cash", "k_min_post_month", None),
                          ("Minimum-cash buffer", "k_buffer", "zar"), ("Shortfall vs buffer", "k_buf_short", "zar"),
                          ("Pre-seed bridge required", "k_bridge", "zar"),
                          ("Nigeria launch", "k_launch_NG", None), ("Kenya launch", "k_launch_KE", None), ("Ghana launch", "k_launch_GH", None),
                          ("Rest of Africa (USD) launch", "k_launch_RoA", None)]:
        r += 1
        _cell(ws, f"A{r}", lab)
        _cell(ws, f"C{r}", f"=Cash_Flow!$B${A.SCALPOS[key]}", F_BOLD, fmt)
    r += 2
    ws.cell(row=r, column=1, value="Changes from v1 (R18.5M model) and why")
    _hdr_row(ws, r, 1, 7)
    for a, b in CHANGES_FROM_V1:
        r += 1
        _cell(ws, f"A{r}", a, F_BOLD, None, None, WRAP)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        _cell(ws, f"B{r}", b, F_BODY, None, None, WRAP)
        ws.row_dimensions[r].height = 44
    r += 2
    ws.cell(row=r, column=1, value="Sheet guide")
    _hdr_row(ws, r, 1, 7)
    guide = [
        ("Assumptions", "All drivers with units and sources; scenario selector (B6); lean mode, pilot, Founding Member, markets, segments; MRR-gated headcount plan."),
        ("Pricing", "ZAR list, NGN/KES/GHS/USD price points, partner wholesale, Founding Member & pilot prices, FX depreciation & repricing, ZAR-parity table."),
        ("Revenue_Build", "Monthly launch gates, FX value factors, trials, marketing cap & payback test, cohorts by market/segment/tier, discounts, revenue."),
        ("Costs", "MRR-gated headcount by role, payroll, COGS (AI, hosting, WhatsApp, processing, support, CS), opex."),
        ("P&L", "Monthly P&L: gross subscriptions -> discounts -> net, other revenue, gross margin, opex, EBITDA, tax, net income."),
        ("Annual_Summary", "FY1-FY5 roll-up: revenue, discounts, margins, EBITDA, cash, customers by segment/market/tier, ARR, ARPA, headcount."),
        ("Cash_Flow", "Working capital, funding (seed only), cash vs R3.0M buffer, runway, break-even, bridge, launch months and the constraint check."),
        ("Unit_Economics", "ARPA, churn, LTV, CAC, LTV:CAC and payback by segment (tier memo); magic number, burn multiple, Rule of 40."),
        ("Use_of_Funds", "R25M allocation and reconciliation against modelled spend in the 18 and 24 months after close."),
        ("Sensitivity", "Scenario comparison, FX shock, discount options, pilot conversion, Conservative shortfall & fix, churn x volume grids (build-time values)."),
        ("Market_Sizing", "TAM / SAM / SOM (estimates) linked to modelled FY5 workspaces."),
    ]
    for name, desc in guide:
        r += 1
        _cell(ws, f"A{r}", name, F_BOLD)
        _cell(ws, f"B{r}", desc)
    r += 2
    ws.cell(row=r, column=1, value="Colour legend")
    _hdr_row(ws, r, 1, 7)
    r += 1
    _input(ws, f"A{r}", "1,234", "text")
    _cell(ws, f"B{r}", "Input: blue font on light yellow. Change these.")
    r += 1
    _cell(ws, f"A{r}", "=1+1", F_BODY)
    _cell(ws, f"B{r}", "Formula: black font. Do not overwrite.")
    r += 1
    _cell(ws, f"A{r}", "5,678", F_BUILD, None, FILL_BUILD)
    _cell(ws, f"B{r}", "Value computed at build time by the Python mirror (Sensitivity sheet only).")
    r += 1
    _cell(ws, f"A{r}", "[[CONFIRM]]", F_FLAG)
    _cell(ws, f"B{r}", "Red source note: founder confirmation or decision needed.")
    r += 2
    _cell(ws, f"A{r}", "Not in any scenario", F_BOLD)
    _cell(ws, f"B{r}", "Campaign Financing (upside only). Partner referral commission (undecided). Series A (optional acceleration only; switched off).")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 18
    for colL in "CDEFG":
        ws.column_dimensions[colL].width = 17
    ws.sheet_view.showGridLines = False


def build_workbook(p, sens):
    layout()
    wb = Workbook()
    wb.active.title = "Cover"
    write_assumptions(wb)
    write_pricing(wb)
    write_monthly(wb)
    write_annual(wb, "Annual_Summary", "FY1-FY5 (Oct-Sep) roll-up of the monthly model, selected scenario.")
    write_annual(wb, "Unit_Economics", "Unit economics by segment and year. LTV uses software gross margin and segment churn.")
    write_use_of_funds(wb, p)
    write_sensitivity(wb, sens)
    write_market_sizing(wb)
    write_cover(wb)
    order = ["Cover", "Assumptions", "Pricing", "Revenue_Build", "Costs", "P&L", "Annual_Summary", "Cash_Flow",
             "Unit_Economics", "Use_of_Funds", "Sensitivity", "Market_Sizing"]
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = ORANGE if ws.title in ("Cover", "Assumptions", "Pricing") else SLATE
    return wb
