"""FluxMuse proposal document kit (python-docx).

House style: A4 portrait, Poppins headings / Inter body (Arial fallback via fontTable altName),
Flux Orange fills with Ink text (never white on orange), Deep Orange for accent text,
[[PLACEHOLDER]] fields bold on a light-yellow shading.
"""
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Cm, RGBColor, Emu
from PIL import Image

HERE = Path(__file__).resolve().parent
GTM = HERE.parent.parent  # docs/go-to-market
ASSETS = GTM / "assets"

INK = "20242B"
ORANGE = "FF6A00"
DEEP = "C24E00"
TINT = "FFF4EB"
SLATE = "37474F"
MIST = "F3F4F6"
GREY = "5B6570"
RULE = "D9DDE2"
PH_FILL = "FFF1A8"  # light yellow for fill-in fields
HEAD_FONT = "Poppins"
BODY_FONT = "Inter"
CONTENT_W = 17.0  # cm (A4 21cm - 2 x 2cm margins)

PH_RE = re.compile(r"(\[\[.*?\]\])")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

# ---------------------------------------------------------------- low-level XML helpers

PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
             "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
             "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
             "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar",
               "tblLook", "tblCaption", "tblDescription"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
              "tblHeader", "tblCellSpacing", "jc", "hidden"]


def put(parent, child, order):
    """Insert/replace child in parent respecting the schema sequence `order`."""
    name = child.tag.split("}")[1]
    for old in parent.findall(qn("w:" + name)):
        parent.remove(old)
    idx = order.index(name)
    for i, el in enumerate(list(parent)):
        tag = el.tag.split("}")[1] if "}" in el.tag else el.tag
        if tag in order and order.index(tag) > idx:
            el.addprevious(child)
            return child
    parent.append(child)
    return child


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def shd(fill):
    return el("w:shd", val="clear", color="auto", fill=fill)


def border_el(tag, sides, sz=4, color=RULE, val="single", space=0):
    b = OxmlElement(tag)
    for s in sides:
        b.append(el("w:" + s, val=val, sz=sz, space=space, color=color))
    return b


def set_fonts(rpr, font):
    put(rpr, el("w:rFonts", ascii=font, hAnsi=font, cs=font, eastAsia=font), RPR_ORDER)


# ---------------------------------------------------------------- images

_img_cache = {}


def img_stream(path, max_w=2400):
    path = Path(path)
    key = (str(path), max_w, path.stat().st_mtime)
    if key not in _img_cache:
        im = Image.open(path)
        if im.width > max_w:
            im = im.resize((max_w, round(im.height * max_w / im.width)), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, "PNG", optimize=True)
        _img_cache[key] = buf.getvalue()
    return io.BytesIO(_img_cache[key])


# ---------------------------------------------------------------- document


class FMDoc:
    def __init__(self, sample=False, client_ph="[[CLIENT NAME]]"):
        self.doc = Document()
        self.sample = sample
        self.client_ph = client_ph
        self.fig_n = 0
        self.tab_n = 0
        self.toc_para = None
        self.h1s = []
        self.images_used = []
        self._setup_page()
        self._setup_styles()
        self._setup_fonts_table()
        self._header_footer()

    # ---- setup
    def _setup_page(self):
        s = self.doc.sections[0]
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin, s.bottom_margin = Cm(2.4), Cm(2.2)
        s.header_distance, s.footer_distance = Cm(1.0), Cm(1.0)
        s.different_first_page_header_footer = True
        props = self.doc.core_properties
        props.author = "Fluxmuse Pty Ltd"
        props.last_modified_by = "Fluxmuse Pty Ltd"

    def _style_font(self, style, font, size=None, color=INK, bold=None, italic=None):
        rpr = style._element.get_or_add_rPr()
        set_fonts(rpr, font)
        style.font.color.rgb = RGBColor.from_string(color)
        # remove theme colour attributes inherited from the template
        c = rpr.find(qn("w:color"))
        for a in ("themeColor", "themeShade", "themeTint"):
            if c is not None and c.get(qn("w:" + a)) is not None:
                del c.attrib[qn("w:" + a)]
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if italic is not None:
            style.font.italic = italic

    def _setup_styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        self._style_font(normal, BODY_FONT, 10.5)
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        # docDefaults fonts too, so tables/lists inherit Inter
        rpr_default = self.doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        set_fonts(rpr_default, BODY_FONT)

        specs = {1: (18, INK, 20, 8), 2: (13, INK, 14, 4), 3: (11, DEEP, 10, 3)}
        for lvl, (size, color, before, after) in specs.items():
            h = st["Heading %d" % lvl]
            self._style_font(h, HEAD_FONT, size, color, bold=True, italic=False)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.05
            if lvl == 1:
                ppr = h._element.get_or_add_pPr()
                put(ppr, border_el("w:pBdr", ["bottom"], sz=12, color=ORANGE, space=4), PPR_ORDER)
        for name in ("Heading 1 Char", "Heading 2 Char", "Heading 3 Char"):
            try:
                cs = st[name]
                self._style_font(cs, HEAD_FONT, None, INK)
            except KeyError:
                pass
        cap = st["Caption"]
        self._style_font(cap, BODY_FONT, 8.5, GREY, bold=False, italic=False)
        cap.paragraph_format.space_before = Pt(3)
        cap.paragraph_format.space_after = Pt(10)
        for name in ("List Bullet", "List Number", "List Bullet 2"):
            s = st[name]
            s.paragraph_format.space_after = Pt(3)
            s.paragraph_format.line_spacing = 1.1
            ppr = s._element.get_or_add_pPr()
            cs = ppr.find(qn("w:contextualSpacing"))
            if cs is not None:
                ppr.remove(cs)
        title = st["Title"]
        self._style_font(title, HEAD_FONT, 30, INK, bold=True)
        ppr = title._element.get_or_add_pPr()
        b = ppr.find(qn("w:pBdr"))
        if b is not None:
            ppr.remove(b)

    def _setup_fonts_table(self):
        ft = None
        for rel in self.doc.part.rels.values():
            if rel.reltype.endswith("/fontTable"):
                ft = rel.target_part
        if ft is None:
            return
        root = parse_xml(ft.blob)
        for name in (HEAD_FONT, BODY_FONT):
            f = parse_xml(
                f'<w:font {nsdecls("w")} w:name="{name}"><w:altName w:val="Arial"/>'
                f'<w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>')
            root.append(f)
        from lxml import etree
        ft._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def _field(self, p, instr, size=8, color=GREY):
        fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr=" {instr} "/>')
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        set_fonts(rpr, BODY_FONT)
        put(rpr, el("w:color", val=color), RPR_ORDER)
        put(rpr, el("w:sz", val=int(size * 2)), RPR_ORDER)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)

    def _footer_para(self, footer):
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ppr = p._p.get_or_add_pPr()
        put(ppr, border_el("w:pBdr", ["top"], sz=4, color=RULE, space=6), PPR_ORDER)
        self.rich(p, f"Fluxmuse Pty Ltd · fluxmuse.ai · Confidential proposal for {self.client_ph} · Page ",
                  size=8, color=GREY)
        self._field(p, "PAGE")
        self.rich(p, " of ", size=8, color=GREY)
        self._field(p, "NUMPAGES")

    def _header_para(self, header):
        p = header.paragraphs[0]
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.tab_stops.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
        ppr = p._p.get_or_add_pPr()
        put(ppr, border_el("w:pBdr", ["bottom"], sz=12, color=ORANGE, space=4), PPR_ORDER)
        r = p.add_run()
        r.add_picture(img_stream(ASSETS / "brand/fluxmuse-icon-512.png", 256), height=Cm(0.8))
        self._docpr(r, "FluxMuse icon")
        self._run(p, "  FLUX", font=HEAD_FONT, size=12, color=DEEP, bold=True)
        self._run(p, "MUSE", font=HEAD_FONT, size=12, color=SLATE, bold=True)
        if self.sample:
            self._run(p, "\tSAMPLE: ILLUSTRATIVE · ", size=8, color=DEEP, bold=True)
            self.rich(p, "Proposal [[FM-PRO-YYYY-###]]", size=8, color=GREY)
        else:
            self.rich(p, "\tProposal [[FM-PRO-YYYY-###]]", size=8, color=GREY)

    def _header_footer(self):
        s = self.doc.sections[0]
        self._header_para(s.header)
        self._footer_para(s.footer)
        self._footer_para(s.first_page_footer)
        s.first_page_header.paragraphs[0].text = ""

    def _docpr(self, run, descr):
        for dp in run._r.iter(qn("wp:docPr")):
            dp.set("descr", descr)

    # ---- inline text
    def _run(self, p, text, font=None, size=None, color=None, bold=None, italic=None, ph=False):
        r = p.add_run(text)
        rpr = r._r.get_or_add_rPr()
        if font:
            set_fonts(rpr, font)
        if bold is not None:
            r.bold = bold
        if italic is not None:
            r.italic = italic
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if size:
            r.font.size = Pt(size)
        if ph:
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(INK)
            put(rpr, shd(PH_FILL), RPR_ORDER)
        return r

    def rich(self, p, text, font=None, size=None, color=None, bold=None, italic=None):
        """Add text with **bold** and [[PLACEHOLDER]] support."""
        parts = BOLD_RE.split(text)
        for i, part in enumerate(parts):
            if not part:
                continue
            b = True if i % 2 == 1 else bold
            for seg in PH_RE.split(part):
                if not seg:
                    continue
                if PH_RE.fullmatch(seg):
                    self._run(p, seg, font=font, size=size, italic=italic, ph=True)
                else:
                    self._run(p, seg, font=font, size=size, color=color, bold=b, italic=italic)
        return p

    # ---- blocks
    def para(self, text="", size=None, color=None, bold=None, italic=None, align=None, after=None,
             before=None, keep=False, style=None, font=None):
        p = self.doc.add_paragraph(style=style)
        self.rich(p, text, size=size, color=color, bold=bold, italic=italic, font=font)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if before is not None:
            p.paragraph_format.space_before = Pt(before)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def h1(self, text, new_page=False):
        p = self.doc.add_heading(level=1)
        self.rich(p, text)
        if new_page:
            p.paragraph_format.page_break_before = True
        self.h1s.append(text)
        return p

    def h2(self, text):
        p = self.doc.add_heading(level=2)
        self.rich(p, text)
        return p

    def h3(self, text):
        p = self.doc.add_heading(level=3)
        self.rich(p, text)
        return p

    def eyebrow(self, text):
        p = self.para(text.upper(), size=8.5, color=DEEP, bold=True, font=HEAD_FONT, after=2, keep=True)
        for r in p.runs:
            put(r._r.get_or_add_rPr(), el("w:spacing", val=30), RPR_ORDER)
        return p

    def bullets(self, items, level=1):
        style = "List Bullet" if level == 1 else "List Bullet 2"
        out = []
        for it in items:
            p = self.doc.add_paragraph(style=style)
            self.rich(p, it)
            out.append(p)
        return out

    def numbered(self, items):
        numbering = self.doc.part.numbering_part.element
        style_num = self.doc.styles["List Number"]._element.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        abstract = None
        for n in numbering.findall(qn("w:num")):
            if n.get(qn("w:numId")) == style_num:
                abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{new_id}"><w:abstractNumId w:val="{abstract}"/>'
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>')
        numbering.append(num)
        out = []
        for it in items:
            p = self.doc.add_paragraph(style="List Number")
            ppr = p._p.get_or_add_pPr()
            numpr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{new_id}"/></w:numPr>')
            put(ppr, numpr, PPR_ORDER)
            self.rich(p, it)
            out.append(p)
        return out

    def checkbox_run(self, p, checked=False):
        sym = "2612" if checked else "2610"
        ch = "☒" if checked else "☐"
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")} xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
            '<w:sdtPr><w:rPr><w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic" w:hint="eastAsia"/></w:rPr>'
            f'<w14:checkbox><w14:checked w14:val="{1 if checked else 0}"/>'
            '<w14:checkedState w14:val="2612" w14:font="MS Gothic"/>'
            '<w14:uncheckedState w14:val="2610" w14:font="MS Gothic"/></w14:checkbox></w:sdtPr>'
            '<w:sdtContent><w:r><w:rPr><w:rFonts w:ascii="MS Gothic" w:eastAsia="MS Gothic" w:hAnsi="MS Gothic" w:hint="eastAsia"/>'
            f'<w:color w:val="{INK}"/></w:rPr><w:t>{ch}</w:t></w:r></w:sdtContent></w:sdt>')
        p._p.append(sdt)

    def checklist(self, items, container=None, keep=False):
        out = []
        for it in items:
            checked = False
            if isinstance(it, tuple):
                it, checked = it
            p = (container.add_paragraph() if container is not None else self.doc.add_paragraph())
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.7))
            if keep:
                p.paragraph_format.keep_with_next = True
            self.checkbox_run(p, checked)
            self.rich(p, "\t" + it)
            out.append(p)
        return out

    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    def spacer(self, pt=6):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(pt)
        return p

    def rule(self, color=ORANGE, sz=18, after=6):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        put(p._p.get_or_add_pPr(), border_el("w:pBdr", ["bottom"], sz=sz, color=color, space=1), PPR_ORDER)
        return p

    # ---- figures
    def figure(self, rel_path, caption, alt, width_cm=CONTENT_W):
        path = ASSETS / rel_path
        self.images_used.append(rel_path)
        self.fig_n += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run()
        r.add_picture(img_stream(path), width=Cm(width_cm))
        self._docpr(r, alt)
        c = self.doc.add_paragraph(style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.rich(c, f"**Figure {self.fig_n}.** {caption}")
        return p

    # ---- tables
    def _table(self, nrows, widths, caption):
        t = self.doc.add_table(rows=nrows, cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        total = sum(widths)
        put(tblPr, el("w:tblW", w=int(Cm(total) / 635), type="dxa"), TBLPR_ORDER)
        put(tblPr, el("w:tblLayout", type="fixed"), TBLPR_ORDER)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", 50), ("left", 100), ("bottom", 50), ("right", 100)):
            mar.append(el("w:" + side, w=v, type="dxa"))
        put(tblPr, mar, TBLPR_ORDER)
        put(tblPr, el("w:tblCaption", val=caption), TBLPR_ORDER)
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(Cm(w) / 635)))
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
        return t

    def cell_fill(self, cell, fill):
        put(cell._tc.get_or_add_tcPr(), shd(fill), TCPR_ORDER)

    def cell_borders(self, cell, spec):
        """spec: dict side -> (sz, color)"""
        b = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            if side in spec:
                sz, color = spec[side]
                b.append(el("w:" + side, val="single" if sz else "nil", sz=sz or 0, space=0, color=color))
        put(cell._tc.get_or_add_tcPr(), b, TCPR_ORDER)

    def cell_text(self, cell, text, size=9.5, bold=None, color=None, align=None, font=None):
        lines = text.split("\n") if isinstance(text, str) else text
        first = True
        for line in lines:
            if first:
                p = cell.paragraphs[0]
                first = False
            else:
                p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if align == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif align == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if line.startswith("[ ] ") or line.startswith("[x] "):
                self.checkbox_run(p, line.startswith("[x]"))
                line = " " + line[4:]
            if line.startswith("- "):
                line = "– " + line[2:]
            self.rich(p, line, size=size, bold=bold, color=color, font=font)
        return cell

    def table(self, headers, rows, widths, caption, size=9.5, first_col_bold=True, zebra=True,
              aligns=None, header_fill=ORANGE, total_rows=(), highlight_rows=(), gantt=False):
        """Data table: header row (Flux Orange fill, Ink bold text) repeats across pages."""
        self.tab_n += 1
        t = self._table(len(rows) + 1, widths, caption)
        tblPr = t._tbl.tblPr
        put(tblPr, border_el("w:tblBorders", ["top", "left", "bottom", "right", "insideH", "insideV"], sz=4, color=RULE), TBLPR_ORDER)
        hdr = t.rows[0]
        trpr = hdr._tr.get_or_add_trPr()
        put(trpr, el("w:cantSplit"), TRPR_ORDER)
        put(trpr, el("w:tblHeader"), TRPR_ORDER)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            self.cell_fill(c, header_fill)
            self.cell_text(c, h, size=size, bold=True, color=INK, font=HEAD_FONT,
                           align=(aligns[i] if aligns else None))
            for hp in c.paragraphs:  # keep the header row with the first data row (no orphaned headers)
                hp.paragraph_format.keep_with_next = True
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            r = t.rows[ri + 1]
            put(r._tr.get_or_add_trPr(), el("w:cantSplit"), TRPR_ORDER)
            for ci, val in enumerate(row):
                c = r.cells[ci]
                if gantt and ci > 0 and val in ("#", "##"):
                    self.cell_fill(c, "FFB27A" if val == "#" else ORANGE)
                    self.cell_text(c, "●", size=7, color=INK, align="center")
                    continue
                bold = (ci == 0 and first_col_bold) or (ri in total_rows)
                self.cell_text(c, val, size=size, bold=True if bold else None,
                               align=(aligns[ci] if aligns else None))
                if ri in total_rows or ri in highlight_rows:
                    self.cell_fill(c, TINT)
                elif zebra and ri % 2 == 1:
                    self.cell_fill(c, "FAFAFB")
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        self.spacer(6)
        return t

    def kv_table(self, rows, widths=(5.0, 12.0), caption="Layout: details", size=10, label_fill=MIST):
        """Two-column label/value layout table (no header row by design)."""
        t = self._table(len(rows), list(widths), caption)
        put(t._tbl.tblPr, border_el("w:tblBorders", ["top", "bottom", "insideH"], sz=4, color=RULE), TBLPR_ORDER)
        for ri, (k, v) in enumerate(rows):
            put(t.rows[ri]._tr.get_or_add_trPr(), el("w:cantSplit"), TRPR_ORDER)
            a, b = t.rows[ri].cells
            self.cell_fill(a, label_fill)
            self.cell_text(a, k, size=size, bold=True)
            self.cell_text(b, v, size=size)
            a.vertical_alignment = b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        self.spacer(6)
        return t

    def callout(self, title, body, fill=TINT, accent=ORANGE, caption="Layout: callout", title_color=INK):
        """One-cell box with a thick left accent. body: list of str; '- ' = bullet, '[ ] ' = checkbox."""
        t = self._table(1, [CONTENT_W], caption)
        put(t._tbl.tblPr, border_el("w:tblBorders", [], sz=0), TBLPR_ORDER)
        c = t.rows[0].cells[0]
        self.cell_fill(c, fill)
        self.cell_borders(c, {"left": (36, accent), "top": (0, "auto"), "bottom": (0, "auto"), "right": (0, "auto")})
        mar = OxmlElement("w:tcMar")
        for side, v in (("top", 140), ("left", 220), ("bottom", 120), ("right", 200)):
            mar.append(el("w:" + side, w=v, type="dxa"))
        put(c._tc.get_or_add_tcPr(), mar, TCPR_ORDER)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        self.rich(p, title, font=HEAD_FONT, size=11.5, bold=True, color=title_color)
        for line in body:
            if line.startswith("- "):
                q = c.add_paragraph(style="List Bullet")
                self.rich(q, line[2:], size=9.5)
            elif line.startswith("[ ] "):
                self.checklist([line[4:]], container=c)
                q = c.paragraphs[-1]
                for r in q.runs:
                    r.font.size = Pt(9.5)
            else:
                q = c.add_paragraph()
                q.paragraph_format.space_after = Pt(4)
                self.rich(q, line, size=9.5)
        self.spacer(8)
        return t

    def signature_block(self, left_title, right_title):
        t = self._table(6, [8.3, 0.4, 8.3], "Layout: signature and acceptance block")
        put(t._tbl.tblPr, border_el("w:tblBorders", [], sz=0), TBLPR_ORDER)
        labels = ["Name", "Title / role", "Signature", "Date", "Place"]
        for side, col in ((left_title, 0), (right_title, 2)):
            c = t.rows[0].cells[col]
            self.cell_fill(c, ORANGE)
            self.cell_text(c, side, size=10, bold=True, color=INK, font=HEAD_FONT)
            for i, lab in enumerate(labels):
                cc = t.rows[i + 1].cells[col]
                self.cell_borders(cc, {"bottom": (6, SLATE)})
                self.cell_text(cc, f"{lab}:", size=9, color=GREY)
                if lab == "Signature":
                    t.rows[i + 1].height = Cm(1.4)
                else:
                    t.rows[i + 1].height = Cm(0.9)
                cc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        for ri, row in enumerate(t.rows):
            put(row._tr.get_or_add_trPr(), el("w:cantSplit"), TRPR_ORDER)
            if ri < len(t.rows) - 1:  # keep the whole signature block on one page
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True
        self.spacer(6)
        return t

    # ---- cover and TOC
    def cover(self, eyebrow, title, subtitle, client_label="Prepared for", extra_rows=()):
        self.spacer(10)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run()
        r.add_picture(img_stream(ASSETS / "brand/fluxmuse-icon-512.png", 512), height=Cm(3.2))
        self._docpr(r, "FluxMuse Muse icon: a multicolour low-poly head over a node network")
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        self._run(p, "FLUX", font=HEAD_FONT, size=30, color=DEEP, bold=True)
        self._run(p, "MUSE", font=HEAD_FONT, size=30, color=SLATE, bold=True)
        self.para("AI Marketing & WhatsApp Commerce for Africa", size=11, color=GREY, after=26)
        if self.sample:
            self.callout("SAMPLE: illustrative",
                         ["A pre-filled example for a Gauteng braiding & hair studio. It is not a real client, "
                          "and every number in it is a target (goal), not a result. Replace the [[PLACEHOLDERS]] "
                          "before using any wording with a real prospect."],
                         fill="FFF1A8", accent=DEEP, caption="Layout: sample banner")
        self.eyebrow(eyebrow)
        p = self.doc.add_paragraph(style="Title")
        p.paragraph_format.space_after = Pt(4)
        self.rich(p, title)
        self.para(subtitle, size=13, color=SLATE, after=14)
        self.rule(ORANGE, 24, after=10)
        rows = [(client_label, self.client_ph),
                ("Date", "[[DD MONTH YYYY]]"),
                ("Proposal number", "[[FM-PRO-YYYY-###]]"),
                ("Prepared by", "[[NAME]], [[TITLE]], Fluxmuse Pty Ltd\n[[EMAIL]] · [[PHONE / WHATSAPP]]"),
                ("Valid until", "[[DATE + 30 days]]")] + list(extra_rows)
        self.kv_table(rows, caption="Layout: proposal details")
        self.para(f"Confidential. Prepared for {self.client_ph} only. Prices in South African rand (ZAR) and "
                  "exclusive of VAT unless stated [[CONFIRM VAT TREATMENT]].", size=8.5, color=GREY, before=10)
        self.page_break()

    def toc(self):
        p = self.para("Contents", font=HEAD_FONT, size=16, bold=True, after=6, keep=True)
        self.rule(ORANGE, 12, after=4)
        self.toc_para = self.doc.add_paragraph()
        self.para("Word builds page numbers when you open the file (choose Yes to update fields), or right-click "
                  "the list and choose Update Field.", size=8.5, color=GREY, italic=True, after=10)

    def _fill_toc(self):
        if self.toc_para is None:
            return
        anchor = self.toc_para
        instr = 'TOC \\o "%s" \\h \\z \\u' % getattr(self, "toc_levels", "1-1")
        entries = [re.sub(r"\*\*|\[\[|\]\]", "", h) for h in self.h1s] or ["(update field)"]
        paras = [anchor]
        for _ in entries[1:]:
            newp = OxmlElement("w:p")
            paras[-1]._p.addnext(newp)
            from docx.text.paragraph import Paragraph
            paras.append(Paragraph(newp, anchor._parent))
        for i, (para, text) in enumerate(zip(paras, entries)):
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
            if i == 0:
                para._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
                r = parse_xml(f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>')
                para._p.append(r)
                para._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'))
            self._run(para, text, size=10.5)
            if i == len(entries) - 1:
                para._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))
        # ask Word to refresh fields (TOC page numbers) on open
        settings = self.doc.settings.element
        upd = el("w:updateFields", val="true")
        after_names = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids"]
        for child in list(settings):
            if child.tag.split("}")[1] in after_names:
                child.addprevious(upd)
                break
        else:
            settings.append(upd)

    def save(self, path, title, subject):
        self._fill_toc()
        # python-docx's default template has <w:zoom w:val="bestFit"/> without the schema-required percent
        z = self.doc.settings.element.find(qn("w:zoom"))
        if z is not None and z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
        cp = self.doc.core_properties
        cp.title = title
        cp.subject = subject
        cp.keywords = "FluxMuse, proposal, template"
        self.doc.save(path)
        return path
