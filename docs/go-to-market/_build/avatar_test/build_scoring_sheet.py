"""Builds docs/go-to-market/08_Flux_Loop/Avatar_Test_Run/Avatar_Test_Scoring.xlsx.

The scoring sheet for the stage 0 avatar video test (../../08_Flux_Loop/
Avatar_Video_Test_Kit.md): brands, one row per video, the voice check,
tracking links, and the six decision gates as live formulas.

Usage:
    python3 docs/go-to-market/_build/avatar_test/build_scoring_sheet.py [--sample OUT.xlsx]

--sample writes a copy filled with made-up test data to OUT.xlsx instead (for
checking the gate formulas; never use it as a result).
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.normpath(os.path.join(HERE, "../../08_Flux_Loop/Avatar_Test_Run/Avatar_Test_Scoring.xlsx"))

ORANGE = "FF6A00"
INK = "20242B"
MIST = "F3F4F6"
TINT = "FFF4EB"
HEAD = Font(name="Inter", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=INK)
INPUT_FILL = PatternFill("solid", fgColor=TINT)
CALC_FILL = PatternFill("solid", fgColor=MIST)
TITLE = Font(name="Poppins", bold=True, size=14, color=INK)
NOTE = Font(name="Inter", italic=True, size=9, color="555555")

BRANDS = [f"{i:02d}" for i in range(0, 11)]  # 00 = FluxMuse dry run, 01-10 = test businesses
FORMATS = ["A", "B", "C"]
# Posting orders rotate across businesses so no format always goes first.
POST_ORDERS = ["A, B, C", "B, C, A", "C, A, B"]
# The order the three videos are labelled 1, 2, 3 for the owner review.
REVIEW_ORDERS = ["B, A, C", "C, B, A", "A, C, B", "B, C, A", "C, A, B", "A, B, C"]

FIRST_VID_ROW = 4
LAST_VID_ROW = FIRST_VID_ROW + len(BRANDS) * len(FORMATS) - 1
FIRST_BRAND_ROW = 4
LAST_BRAND_ROW = FIRST_BRAND_ROW + len(BRANDS) - 1


def header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = NOTE


def yn(ws, rng):
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def brands_sheet(wb):
    ws = wb.active
    ws.title = "Brands"
    title(ws, "Businesses in the test", "00 is FluxMuse (dry run, excluded from the gates). Orange cells are yours to fill in; grey cells are calculated.")
    cols = ["Code", "Business", "Segment (Solo/SME)", "Category", "Main language", "Owner self-films? (Y/N)",
            "WhatsApp number (intl, digits only)", "Consent date", "Avatar #", "Review order (labels 1,2,3)",
            "Posting order", "Owner reports bad customer reaction? (Y/N)", "AI / fake / scam comments (total)"]
    header(ws, 3, cols, [7, 24, 12, 16, 14, 12, 20, 13, 9, 16, 13, 18, 16])
    for i, code in enumerate(BRANDS):
        r = FIRST_BRAND_ROW + i
        ws.cell(row=r, column=1, value=code)
        if code == "00":
            vals = ["FluxMuse (dry run)", "Own brand", "Software", "English", "Y", "27742426065", None, 1]
            for c, v in enumerate(vals, start=2):
                ws.cell(row=r, column=c, value=v)
            ws.cell(row=r, column=11, value="C, A, B")
        else:
            ws.cell(row=r, column=11, value=POST_ORDERS[(i - 1) % 3])
        ws.cell(row=r, column=10, value=REVIEW_ORDERS[i % len(REVIEW_ORDERS)])
        # Each business's three videos sit in fixed rows on Videos, so a plain SUM
        # needs no text matching (engines disagree on comparing "01" to 1).
        v0 = FIRST_VID_ROW + i * len(FORMATS)
        ws.cell(row=r, column=13, value=f'=SUM(Videos!J{v0}:J{v0 + len(FORMATS) - 1})')
        for c in range(2, 13):
            if c not in (10, 11):
                ws.cell(row=r, column=c).fill = INPUT_FILL
        for c in (10, 11, 13):
            ws.cell(row=r, column=c).fill = CALC_FILL
    yn(ws, f"F{FIRST_BRAND_ROW}:F{LAST_BRAND_ROW}")
    yn(ws, f"L{FIRST_BRAND_ROW}:L{LAST_BRAND_ROW}")


def videos_sheet(wb):
    ws = wb.create_sheet("Videos")
    title(ws, "One row per video", "Fill in after the owner review (D), then 72 hours after each post (E-J). Quality pass (K) is for format A only. Count every attempt in L and every rand in M, including failed ones.")
    cols = ["Code", "Business", "Format", "Owner would post? (Y/N)", "Posted? (Y/N)", "Post date & time", "Views",
            "Conversations (code)", "Orders", "AI / fake / scam comments", "Quality pass? (A only, Y/N)",
            "Attempts", "Cost (R, incl. failed)", "Staff minutes", "Tool / model", "Notes (owner's words, what went wrong)",
            "In gates? (1 = test business, 0 = FluxMuse)"]
    header(ws, 3, cols, [9, 9, 8, 12, 10, 16, 9, 13, 9, 13, 13, 9, 13, 10, 20, 40, 12])
    r = FIRST_VID_ROW
    for code in BRANDS:
        for f in FORMATS:
            ws.cell(row=r, column=1, value=f'="AV"&B{r}&C{r}').fill = CALC_FILL
            ws.cell(row=r, column=2, value=code)
            ws.cell(row=r, column=3, value=f)
            for c in range(4, 17):
                ws.cell(row=r, column=c).fill = INPUT_FILL
            if f != "A":
                ws.cell(row=r, column=11).fill = CALC_FILL
            # Numeric, so the gates never depend on how an engine compares the text code "00".
            ws.cell(row=r, column=17, value=0 if code == "00" else 1).fill = CALC_FILL
            r += 1
    yn(ws, f"D{FIRST_VID_ROW}:E{LAST_VID_ROW}")
    yn(ws, f"K{FIRST_VID_ROW}:K{LAST_VID_ROW}")


def voice_sheet(wb):
    ws = wb.create_sheet("Voice check")
    title(ws, "Voice check (before any video)", "Rated by a fluent speaker of each language: natural / OK / poor / not available. Script: Production_Prompts.md §2.")
    cols = ["Tool / model", "South African English", "isiZulu", "Sesotho", "Rated by (initials)", "Notes"]
    header(ws, 3, cols, [30, 18, 14, 14, 14, 40])
    tools = ["Higgsfield: ElevenLabs v4", "Higgsfield: Text to Speech v2 (MiniMax)", "Higgsfield: Seed Audio",
             "Google Cloud Text-to-Speech", "Veo (speech in video)", "Other:"]
    dv = DataValidation(type="list", formula1='"natural,OK,poor,not available"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, t in enumerate(tools):
        r = 4 + i
        ws.cell(row=r, column=1, value=t)
        for c in range(2, 7):
            ws.cell(row=r, column=c).fill = INPUT_FILL
    dv.add("B4:D9")


def tracking_sheet(wb):
    ws = wb.create_sheet("Tracking")
    title(ws, "Tracking links and captions", "Built from the Brands tab. See Tracking_Codes.md for how to count conversations.")
    cols = ["Code", "Business", "wa.me link", "Caption line", "Fallback (no tappable link)"]
    header(ws, 3, cols, [9, 22, 70, 60, 34])
    r = 4
    for i, code in enumerate(BRANDS):
        br = FIRST_BRAND_ROW + i
        for f in FORMATS:
            ws.cell(row=r, column=1, value=f'="AV"&"{code}"&"{f}"')
            ws.cell(row=r, column=2, value=f"=Brands!B{br}")
            ws.cell(row=r, column=3, value=(
                f'=IF(OR(Brands!G{br}="",Brands!B{br}=""),"(fill in the business name and number on Brands)",'
                f'"https://wa.me/"&Brands!G{br}&"?text="&ENCODEURL("Hi "&Brands!B{br}&"! I saw your video ("&A{r}&")"))'))
            ws.cell(row=r, column=4, value=f'=IF(LEFT(C{r},5)="https","Tap to chat: "&C{r},"")')
            ws.cell(row=r, column=5, value=f'="WhatsApp us \'"&A{r}&"\' to order"')
            for c in range(1, 6):
                ws.cell(row=r, column=c).fill = CALC_FILL
            r += 1


def gates_sheet(wb):
    ws = wb.create_sheet("Gates")
    title(ws, "Decision gates (Avatar_Video_Test_Kit.md §7)", "Calculated from the other tabs, excluding 00 (FluxMuse). NEED DATA means not enough rows are filled in yet.")
    V = lambda col: f"Videos!${col}${FIRST_VID_ROW}:${col}${LAST_VID_ROW}"
    B, C, D, E, H, J, K, L, M, Q = (V(c) for c in "BCDEHJKLMQ")
    BR = lambda col: f"Brands!${col}${FIRST_BRAND_ROW + 1}:${col}${LAST_BRAND_ROW}"  # 01-10 only

    ws["A4"] = "Cost limit per finished format A video (R)"
    ws["A4"].font = Font(name="Inter", bold=True)
    ws["B4"].fill = INPUT_FILL
    ws["C4"] = "Set this before day 1 (the dry run tells you a realistic figure)."
    ws["C4"].font = NOTE

    cols = ["Gate", "Pass if", "Measured", "Result"]
    for i, (c, w) in enumerate(zip(cols, [24, 60, 26, 14]), start=1):
        cell = ws.cell(row=6, column=i, value=c)
        cell.font = HEAD
        cell.fill = HEAD_FILL
        ws.column_dimensions[cell.column_letter].width = w

    reviewed_a = f'(COUNTIFS({C},"A",{Q},1,{D},"Y")+COUNTIFS({C},"A",{Q},1,{D},"N"))'
    posted = lambda f: f'COUNTIFS({C},"{f}",{Q},1,{E},"Y")'
    conv = lambda f: f'SUMIFS({H},{C},"{f}",{Q},1,{E},"Y")'
    scored_a = f'(COUNTIFS({C},"A",{Q},1,{K},"Y")+COUNTIFS({C},"A",{Q},1,{K},"N"))'
    passed_a = f'COUNTIFS({C},"A",{Q},1,{K},"Y",{L},"<=2")'
    good_a = f'COUNTIFS({C},"A",{Q},1,{K},"Y")'
    cost_a = f'SUMIFS({M},{C},"A",{Q},1)'
    voice = "'Voice check'!$C$4:$D$9"

    gates = [
        ("Owners post it", "At least half of the owners who reviewed format A actually posted it (needs at least 6 businesses)",
         f'=IF({reviewed_a}=0,"",TEXT({posted("A")}/{reviewed_a},"0%")&" of "&{reviewed_a})',
         f'=IF({reviewed_a}<6,"NEED DATA",IF({posted("A")}/{reviewed_a}>=0.5,"PASS","FAIL"))'),
        ("It works at least as well", "Across businesses, format A brings at least as many conversations per post as format B",
         f'=IF(OR({posted("A")}=0,{posted("B")}=0),"","A "&TEXT({conv("A")}/{posted("A")},"0.0")&" vs B "&TEXT({conv("B")}/{posted("B")},"0.0")&" per post")',
         f'=IF(OR({posted("A")}=0,{posted("B")}=0),"NEED DATA",IF({conv("A")}/{posted("A")}>={conv("B")}/{posted("B")},"PASS","FAIL"))'),
        ("No trust damage", "No owner reports customers calling it fake or a scam, and no more than 1 such comment per business",
         f'="Owner reports: "&COUNTIF({BR("L")},"Y")&"; most comments for one business: "&MAX({BR("M")})',
         f'=IF(AND(COUNTIF({BR("L")},"Y")=0,COUNTIFS({E},"Y",{Q},1)=0),"NEED DATA",IF(AND(COUNTIF({BR("L")},"Y")=0,MAX({BR("M")})<=1),"PASS","FAIL"))'),
        ("Quality", "At least 80% of format A videos pass the quality bar within 2 attempts",
         f'=IF({scored_a}=0,"",TEXT({passed_a}/{scored_a},"0%")&" of "&{scored_a})',
         f'=IF({scored_a}=0,"NEED DATA",IF({passed_a}/{scored_a}>=0.8,"PASS","FAIL"))'),
        ("Local voice", "At least one tool rated natural in isiZulu or Sesotho by a fluent speaker",
         f'=COUNTIF({voice},"natural")&" natural rating(s)"',
         f'=IF(COUNTIF({voice},"natural")>0,"PASS",IF(COUNTA({voice})=0,"NEED DATA","FAIL"))'),
        ("Cost", "Real cost per finished format A video (all attempts, incl. failed) is at or under the limit in B4",
         f'=IF({good_a}=0,IF({cost_a}>0,"R"&TEXT({cost_a},"0")&" spent, 0 finished",""),"R"&TEXT({cost_a}/{good_a},"0")&" per finished video")',
         f'=IF($B$4="","NEED DATA",IF({good_a}=0,IF({cost_a}>0,"FAIL","NEED DATA"),IF({cost_a}/{good_a}<=$B$4,"PASS","FAIL")))'),
    ]
    for i, (g, rule, meas, res) in enumerate(gates):
        r = 7 + i
        ws.cell(row=r, column=1, value=g).font = Font(name="Inter", bold=True)
        ws.cell(row=r, column=2, value=rule).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3, value=meas).fill = CALC_FILL
        cell = ws.cell(row=r, column=4, value=res)
        cell.fill = CALC_FILL
        cell.font = Font(name="Inter", bold=True)
        ws.row_dimensions[r].height = 30
    ws["A14"] = "What to do with the result: Avatar_Video_Test_Kit.md §7."
    ws["A14"].font = NOTE


def fill_sample(wb):
    """Made-up rows so the gate formulas can be checked. Never a result."""
    b = wb["Brands"]
    for i in range(1, 7):
        r = FIRST_BRAND_ROW + i
        b.cell(row=r, column=2, value=f"Sample {i}")
        b.cell(row=r, column=7, value=f"2771000000{i}")
        b.cell(row=r, column=12, value="N")
    v = wb["Videos"]
    for r in range(FIRST_VID_ROW + 3, FIRST_VID_ROW + 3 + 6 * 3):
        f = v.cell(row=r, column=3).value
        n = (r - FIRST_VID_ROW) // 3
        v.cell(row=r, column=4, value="Y")
        v.cell(row=r, column=5, value="Y" if (f != "A" or n % 2) else "N")
        v.cell(row=r, column=8, value={"A": 3, "B": 2, "C": 4}[f])
        v.cell(row=r, column=10, value=0)
        if f == "A":
            v.cell(row=r, column=11, value="Y")
            v.cell(row=r, column=12, value=2)
            v.cell(row=r, column=13, value=60)
    vc = wb["Voice check"]
    vc["C4"] = "OK"
    vc["D5"] = "natural"
    wb["Gates"]["B4"] = 80


def main():
    wb = Workbook()
    brands_sheet(wb)
    videos_sheet(wb)
    voice_sheet(wb)
    tracking_sheet(wb)
    gates_sheet(wb)
    out = OUT
    if "--sample" in sys.argv:
        fill_sample(wb)
        out = sys.argv[sys.argv.index("--sample") + 1]
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
