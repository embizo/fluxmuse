#!/usr/bin/env python3
"""QA for the FluxMuse pitch decks (no LibreOffice needed).

1. Re-opens every deck with python-pptx: slide counts, speaker notes on every slide (60-120 words),
   forbidden claims, leftover template text, placeholders.
2. Geometry: shapes off-slide, edge margins, overlaps between unrelated groups, font sizes by role,
   body word counts.
3. Exports a layout JSON and renders every slide to PNG in headless Chrome (preview_deck.mjs), measuring
   real text overflow with Poppins/Inter, then writes contact sheets.

    python qa_deck.py [preview_dir]      # default preview dir: ./preview
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
DECKS = HERE.parent.parent / "02_Brand_Pitch_Deck"
EMU_IN = 914400

FORBIDDEN = [r"200\+", r"\b4\.9\b", r"snapscan", r"stitch", r"lifetime", r"for life", r"life of your account",
             r"40\s*[–-]\s*60\s*%", r"guarantee", r"flutterwave", r"payfast", r"m-pesa",
             # a Founding Member price must never be shown for the Agency tier
             r"founding member[^.\n]{0,60}(R5,599|R3,999|R5,599/mo)", r"(R5,599|R3,999)[^.\n]{0,40}founding member"]
LEFTOVER = r"lorem|ipsum|\bTODO\b|\bx{3,}\b|\[insert|click to (add|edit)"
MIN_PT = {"body": 14, "card": 14, "chip": 14, "placeholder": 14, "table": 14, "title": 24, "eyebrow": 11,
          "caption": 10, "footer": 10, "slidenum": 10}
WORD_ROLES = {"body", "card", "chip"}


def inch(v):
    return round(Emu(v) / EMU_IN, 3)


def parse_name(name):
    parts = (name or "").split("|")
    return (parts + ["", "", ""])[:3] if len(parts) >= 2 else ("", "", name)


def shape_text_runs(shape):
    paras = []
    for p in shape.text_frame.paragraphs:
        runs = []
        for r in p.runs:
            f = r.font
            color = None
            try:
                color = str(f.color.rgb) if f.color and f.color.type is not None else None
            except Exception:
                pass
            rpr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
            runs.append({"text": r.text, "size": f.size.pt if f.size else 18, "bold": bool(f.bold),
                         "font": f.name or "Inter", "color": color,
                         "spc": int(rpr.get("spc")) / 100 if rpr is not None and rpr.get("spc") else 0})
        for fld in p._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}fld"):
            if len(runs) >= len(p._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")) + 1:
                continue  # this python-pptx version already returns fields as runs
            runs.append({"text": "".join(fld.itertext()), "size": 10, "bold": False, "font": "Inter", "color": None, "spc": 0})
        algn = p._p.pPr.get("algn") if p._p.pPr is not None else None
        paras.append({"runs": runs, "align": algn or "l",
                      "lsp": p.line_spacing if isinstance(p.line_spacing, float) else None,
                      "space_after": p.space_after.pt if p.space_after is not None else 0})
    return paras


def fill_hex(shape):
    try:
        if shape.fill.type == 1:
            return str(shape.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def analyse(path: Path, pv_dir: Path):
    prs = Presentation(path)
    W, H = inch(prs.slide_width), inch(prs.slide_height)
    issues, layout, placeholders = [], {"file": path.name, "w": W, "h": H, "slides": []}, set()
    img_dir = pv_dir / "img"
    img_dir.mkdir(parents=True, exist_ok=True)
    for idx, slide in enumerate(prs.slides, 1):
        where = f"{path.name} s{idx}"
        # notes
        note = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        nw = len(note.split())
        if nw == 0:
            issues.append((where, "notes", "missing speaker notes"))
        elif not 60 <= nw <= 120:
            issues.append((where, "notes", f"notes {nw} words (target 60-120)"))
        bg = None
        try:
            bg = str(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        sl = {"n": idx, "bg": bg or "FFFFFF", "els": []}
        all_text = [note]
        boxes, body_words = [], 0
        for shp in slide.shapes:
            grp, role, label = parse_name(shp.name)
            x, y, w, h = inch(shp.left), inch(shp.top), inch(shp.width), inch(shp.height)
            el = {"name": shp.name, "role": role, "x": x, "y": y, "w": w, "h": h}
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = shp.image.blob
                hsh = hashlib.sha1(blob).hexdigest()[:16]
                fp = img_dir / f"{hsh}.{shp.image.ext}"
                if not fp.exists():
                    fp.write_bytes(blob)
                geom = shp._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
                el.update(type="pic", src=str(fp), crop=[shp.crop_left, shp.crop_top, shp.crop_right, shp.crop_bottom],
                          rounded=geom is not None and geom.get("prst") == "roundRect")
            elif shp.has_table if hasattr(shp, "has_table") else False:
                tbl = shp.table
                cols = [inch(c.width) for c in tbl.columns]
                rows = []
                for r in tbl.rows:
                    cells = []
                    for c in r.cells:
                        f = None
                        try:
                            f = str(c.fill.fore_color.rgb)
                        except Exception:
                            pass
                        paras = shape_text_runs(c)
                        cells.append({"paras": paras, "fill": f})
                        for p in paras:
                            for rr in p["runs"]:
                                all_text.append(rr["text"])
                                if rr["size"] < MIN_PT["table"]:
                                    issues.append((where, "font", f"table run {rr['size']}pt < 14: {rr['text'][:30]}"))
                    rows.append({"h": inch(r.height), "cells": cells})
                el.update(type="table", cols=cols, rows=rows)
                role = "table"
            elif shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                prst = shp._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
                adj = None
                try:
                    adj = shp.adjustments[0] if len(shp.adjustments) else None
                except Exception:
                    pass
                line = None
                try:
                    if shp.line.fill.type == 1:
                        line = {"color": str(shp.line.color.rgb), "w": shp.line.width.pt,
                                "dash": shp.line.dash_style is not None}
                except Exception:
                    pass
                el.update(type="shape", prst=prst.get("prst") if prst is not None else "rect", adj=adj,
                          fill=fill_hex(shp), line=line)
            if shp.has_text_frame and el.get("type") != "pic":
                paras = shape_text_runs(shp)
                txt = "".join(r["text"] for p in paras for r in p["runs"])
                if txt.strip():
                    anchor = shp.text_frame._txBody.bodyPr.get("anchor") or "t"
                    el.update(type="text" if el.get("type") != "shape" else "shape", paras=paras, anchor=anchor)
                    all_text.append(txt)
                    placeholders.update(re.findall(r"\[\[[^\]]+\]\]", txt))
                    for p in paras:
                        for r in p["runs"]:
                            mn = MIN_PT.get(role)
                            if mn and r["size"] < mn and r["text"].strip():
                                issues.append((where, "font", f"{role} run {r['size']}pt < {mn}: {r['text'][:40]}"))
                    if role in WORD_ROLES:
                        body_words += len(txt.split())
            sl["els"].append(el)
            # geometry checks
            if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
                issues.append((where, "offslide", f"{shp.name} at ({x},{y},{w},{h})"))
            if role not in ("bg", "deco", "footer", "slidenum", "") and (x < 0.4 or x + w > W - 0.4 or y < 0.35 or y + h > H - 0.15):
                issues.append((where, "margin", f"{shp.name} near edge ({x},{y},{w},{h})"))
            if role not in ("bg", "deco"):
                boxes.append((grp, role, shp.name, x, y, w, h))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                if a[0] and a[0] == b[0]:
                    continue
                # a card that fully contains the other element is its container, not a collision
                def contains(o, i_):
                    return (o[1] == "card" and o[3] <= i_[3] + 0.01 and o[4] <= i_[4] + 0.01
                            and o[3] + o[5] >= i_[3] + i_[5] - 0.01 and o[4] + o[6] >= i_[4] + i_[6] - 0.01)
                if contains(a, b) or contains(b, a):
                    continue
                ix = min(a[3] + a[5], b[3] + b[5]) - max(a[3], b[3])
                iy = min(a[4] + a[6], b[4] + b[6]) - max(a[4], b[4])
                if ix > 0.02 and iy > 0.02:
                    issues.append((where, "overlap", f"{a[2]} x {b[2]} ({ix:.2f}x{iy:.2f} in)"))
        # appendix reference slides (FAQ, price tables) are exempt from the ~30-word body limit
        is_appendix = any(e["role"] == "eyebrow" and "".join(r["text"] for p in e.get("paras", []) for r in p["runs"])
                          .upper().startswith("APPENDIX") for e in sl["els"])
        if body_words > 35 and not is_appendix:
            issues.append((where, "words", f"{body_words} body words (target <= ~30)"))
        joined = "\n".join(all_text)
        for pat in FORBIDDEN:
            for m in re.finditer(pat, joined, flags=re.I):
                ctx = joined[max(0, m.start() - 40): m.end() + 40].replace("\n", " ")
                issues.append((where, "FORBIDDEN", f"'{m.group(0)}' in: ...{ctx}..."))
        if re.search(LEFTOVER, joined, flags=re.I):
            issues.append((where, "leftover", re.search(LEFTOVER, joined, flags=re.I).group(0)))
        layout["slides"].append(sl)
    return prs, issues, layout, placeholders


def contact_sheets(png_dir: Path, prefix: str, per=6):
    pngs = sorted(png_dir.glob("slide-*.png"))
    sheets = []
    for k in range(0, len(pngs), per):
        chunk = pngs[k:k + per]
        tw, th = 800, 450
        cols = 2
        rows = (len(chunk) + 1) // 2
        sheet = Image.new("RGB", (cols * tw + 30, rows * th + (rows + 1) * 10), (120, 120, 120))
        for i, p in enumerate(chunk):
            im = Image.open(p).convert("RGB").resize((tw, th))
            sheet.paste(im, (10 + (i % 2) * (tw + 10), 10 + (i // 2) * (th + 10)))
        out = png_dir.parent / f"{prefix}-sheet-{k // per + 1}.jpg"
        sheet.save(out, quality=85)
        sheets.append(out)
    return sheets


def main():
    pv_root = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "preview"
    pv_root.mkdir(parents=True, exist_ok=True)
    total = 0
    all_placeholders = set()
    for path in sorted(DECKS.glob("*.pptx")):
        stem = path.stem
        pv = pv_root / stem
        prs, issues, layout, ph = analyse(path, pv)
        all_placeholders |= ph
        lj = pv / "layout.json"
        lj.write_text(json.dumps(layout))
        # visual render + overflow measurement
        env = dict(os.environ)
        r = subprocess.run(["node", str(HERE / "preview_deck.mjs"), str(lj), str(pv / "png")], capture_output=True,
                           text=True, env=env)
        if r.returncode != 0:
            issues.append((path.name, "preview", r.stderr[-400:]))
        else:
            rep = json.loads((pv / "png" / "overflow.json").read_text())
            for item in rep:
                issues.append((f"{path.name} s{item['slide']}", "overflow", item["msg"]))
            contact_sheets(pv / "png", stem)
        notes_ok = all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
        print(f"\n== {path.name}: {len(prs.slides)} slides, notes on every slide: {notes_ok}")
        for w, kind, msg in issues:
            print(f"  [{kind}] {w}: {msg}")
        total += len(issues)
    print("\nPlaceholders found:", ", ".join(sorted(all_placeholders)))
    print(f"\nTotal findings: {total}")


if __name__ == "__main__":
    main()
