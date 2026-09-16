#!/usr/bin/env python3
"""Build the FluxMuse brand pitch decks (master + solo, SME and partner variants).

Every image is embedded from docs/go-to-market/assets/ at build time, so re-running this
script picks up re-rendered infographics. Facts come from 00_FACTS_AND_ASSUMPTIONS.md.

    python build_deck.py            # needs python-pptx + Pillow
    python qa_deck.py               # structural + visual QA (see that file)

Shape names follow "group|role|label" so qa_deck.py can check font sizes, overlaps and
word counts by role. Roles: bg, deco, title, eyebrow, body, card, chip, caption, footer,
slidenum, img, icon, table, placeholder.
"""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
GTM = HERE.parent.parent  # docs/go-to-market
ASSETS = GTM / "assets"
ICONS = HERE / "icons"
OUT = GTM / "02_Brand_Pitch_Deck"
V3_MARKER = ASSETS / ".revision-v3-done"

# ---------------------------------------------------------------- brand tokens
ORANGE, GLOW, TINT = "FF6A00", "FF8533", "FFF4EB"
DEEP = "C24E00"  # orange text on white (AA)
SLATE, INK, NIGHT, PAPER, MIST = "37474F", "20242B", "0F1419", "FFFFFF", "F3F4F6"
NIGHT_CARD = "1B232C"
MUTED = "5B6670"  # captions on white (>= 5:1)
MUTED_DARK = "A7B0B8"  # secondary text on Night
LINE = "E3E6EA"
HEAD, BODY = "Poppins", "Inter"

SW, SH = 13.333, 7.5
MX = 0.6  # side margin
CW = SW - 2 * MX  # content width 12.133

USED_IMAGES: dict[str, set] = {}
_CUR_DECK = [""]
_gid = [0]


def gid(prefix="g"):
    _gid[0] += 1
    return f"{prefix}{_gid[0]}"


# ---------------------------------------------------------------- primitives
def rgb(h):
    return RGBColor.from_string(h)


def set_bg(slide, colour):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(colour)


def text(slide, x, y, w, h, paras, *, size=16, font=BODY, color=INK, bold=False, align="l",
         anchor="t", role="body", group=None, label="", lsp=None, spc=None, space_after=0):
    """paras: str | list[str | list[(text, opts)] | dict(runs=..., size=, align=, space_after=)]"""
    shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = f"{group or gid()}|{role}|{label}"
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        popts = {}
        if isinstance(para, dict):
            popts = para
            runs = para["runs"]
        else:
            runs = para
        if isinstance(runs, str):
            runs = [(runs, {})]
        runs = [(r, {}) if isinstance(r, str) else r for r in runs]
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[popts.get("align", align)]
        if popts.get("lsp", lsp):
            p.line_spacing = popts.get("lsp", lsp)
        sa = popts.get("space_after", space_after)
        if sa:
            p.space_after = Pt(sa)
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", popts.get("size", size)))
            f.name = o.get("font", popts.get("font", font))
            f.bold = o.get("bold", popts.get("bold", bold))
            f.color.rgb = rgb(o.get("color", popts.get("color", color)))
            s = o.get("spc", spc)
            if s:
                r._r.get_or_add_rPr().set("spc", str(s))
    return shp


def box(slide, x, y, w, h, *, fill=None, line=None, lw=1.0, shape="rect", radius=0.14, dash=False,
        role="card", group=None, label=""):
    st = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL,
          "down": MSO_SHAPE.DOWN_ARROW, "right": MSO_SHAPE.RIGHT_ARROW}[shape]
    s = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = f"{group or gid()}|{role}|{label}"
    if shape == "round":
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def pic(slide, rel, x, y, w, h, *, crop=(0, 0, 0, 0), align="c", alt="", group=None, role="img",
        rounded=False, base=ASSETS):
    """Fit the cropped image inside the (x, y, w, h) box. align: c | t | l | tl."""
    path = base / rel
    W, H = Image.open(path).size
    l, t, r, b = crop
    cw, ch = W * (1 - l - r), H * (1 - t - b)
    sc = min(w / cw, h / ch)
    pw, ph = cw * sc, ch * sc
    px = x if "l" in align else x + (w - pw) / 2
    py = y if "t" in align else y + (h - ph) / 2
    p = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = l, t, r, b
    p.name = f"{group or gid()}|{role}|{rel}"
    p._element.nvPicPr.cNvPr.set("descr", alt or Path(rel).stem.replace("-", " "))
    if rounded:
        p.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    if base == ASSETS:
        USED_IMAGES.setdefault(_CUR_DECK[0], set()).add(rel)
    return px, py, pw, ph


def icon(slide, name, colour, cx, cy, d, *, circle=None, ring=None, group=None, scale=0.56):
    """Icon centred at (cx, cy) inside an optional circle of diameter d."""
    g = group or gid()
    if circle or ring:
        box(slide, cx - d / 2, cy - d / 2, d, d, fill=circle, line=ring, lw=2, shape="oval", role="deco", group=g)
    s = d * scale
    pic(slide, f"{name}_{colour}.png", cx - s / 2, cy - s / 2, s, s, base=ICONS, role="icon", group=g, alt="")
    return g


def cell_lines(cell, bottom=LINE, width_pt=1.0):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        if tag == "a:lnB" and bottom:
            ln.set("w", str(int(width_pt * 12700)))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", bottom)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)


def table(slide, x, y, col_w, row_h, rows, *, size=15, group=None, label="table"):
    """rows: list of list of cell (str | dict(text, bold, color, fill, align, size, font))."""
    nr, nc = len(rows), len(col_w)
    heights = row_h if isinstance(row_h, list) else [row_h] * nr
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(heights)))
    gf.name = f"{group or gid()}|table|{label}"
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    tbl = gf.table
    for i, cwid in enumerate(col_w):
        tbl.columns[i].width = Inches(cwid)
    for r, hh in enumerate(heights):
        tbl.rows[r].height = Inches(hh)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            d = val if isinstance(val, dict) else {"text": val}
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(d.get("fill", PAPER))
            cell_lines(cell, bottom=d.get("line", LINE))
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[d.get("align", "l" if c == 0 else "c")]
            run = p.add_run()
            run.text = d["text"]
            f = run.font
            f.size = Pt(d.get("size", size))
            f.name = d.get("font", BODY)
            f.bold = d.get("bold", False)
            f.color.rgb = rgb(d.get("color", INK))
    return gf


def slide_number_box(slide, n, dark):
    shp = slide.shapes.add_textbox(Inches(SW - MX - 1.0), Inches(7.04), Inches(1.0), Inches(0.26))
    shp.name = "chrome|slidenum|"
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"))
    rPr.set("lang", "en-ZA")
    rPr.set("sz", "1000")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", MUTED_DARK if dark else MUTED)
    etree.SubElement(rPr, qn("a:latin")).set("typeface", BODY)
    etree.SubElement(fld, qn("a:t")).text = str(n)


def chrome(slide, n, dark):
    text(slide, MX, 7.04, 4.0, 0.26, "FluxMuse · fluxmuse.ai", size=10, color=MUTED_DARK if dark else MUTED,
         role="footer", group="chrome")
    slide_number_box(slide, n, dark)


def new_slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, NIGHT if dark else PAPER)
    s._fm_dark = dark
    return s


def header(slide, eyebrow, title, *, dark=False, title_size=32, y=0.42, w=CW):
    g = gid("hdr")
    text(slide, MX, y, w, 0.26, eyebrow.upper(), size=12, bold=True, color=ORANGE if dark else DEEP,
         role="eyebrow", group=g, spc=150)
    text(slide, MX, y + 0.28, w, 0.78, title, size=title_size, font=HEAD, bold=True,
         color=PAPER if dark else INK, role="title", group=g, anchor="t")


def notes(slide, note, v):
    t = note.get(v, note["default"]) if isinstance(note, dict) else note
    slide.notes_slide.notes_text_frame.text = " ".join(t.split())


def hero(slide, rel, crop, alt, box_=(MX, 1.52, CW, 5.4), align="c"):
    return pic(slide, rel, *box_, crop=crop, alt=alt, align=align)


# ---------------------------------------------------------------- image crops (l, t, r, b)
# Infographics share one template: eyebrow + title (+ subtitle) in the top ~19% of the canvas,
# which the slide replaces with a native, editable title.
TOP = (0, 0.195, 0, 0)
IMG = {
    "problem": ("infographics/problem-solution.png", (0.03, 0.2, 0.03, 0.15)),
    "hero": ("screenshots/composites/hero-laptop-phone_night.png", (0.02, 0.09, 0.06, 0.02)),
    "how": ("infographics/how-fluxmuse-works.png", (0, 0.2, 0, 0)),
    "roster": ("infographics/ai-agent-roster.png", (0, 0.19, 0, 0.02)),
    "flow": ("infographics/whatsapp-commerce-flow.png", (0.035, 0.272, 0.035, 0.37)),
    "phones": ("screenshots/composites/checkout-in-chat-3-phones.png", (0.12, 0.205, 0.12, 0.02)),
    "hub": ("infographics/omnichannel-hub.png", TOP),
    "stack": ("infographics/platform-stack.png", TOP),
    "map": ("infographics/payment-coverage-map.png", (0, 0.205, 0, 0)),
    "solo": ("infographics/segment-solo.png", (0, 0.2, 0, 0)),
    "sme": ("infographics/segment-sme.png", (0, 0.2, 0, 0)),
    "agency": ("infographics/segment-agency.png", (0, 0.2, 0, 0)),
    "partner": ("infographics/agency-partner-model.png", TOP),
    "laptop": ("screenshots/composites/laptop-home.png", (0.03, 0.08, 0.03, 0.08)),
    "pricing_shot": ("screenshots/composites/browser-pricing.png", (0.09, 0.375, 0.09, 0.39)),
    "ph1": ("screenshots/composites/phone-demo-step-01-catalog.png", (0.06, 0.03, 0.06, 0.03)),
    "ph3": ("screenshots/composites/phone-demo-step-03-cart.png", (0.06, 0.03, 0.06, 0.03)),
    "ph7": ("screenshots/composites/phone-demo-step-07-confirmation.png", (0.06, 0.03, 0.06, 0.03)),
    "journey": ("infographics/brand-engagement-journey.png", (0, 0.22, 0, 0)),
    "pilot": ("infographics/gauteng-pilot.png", (0, 0.2, 0, 0)),
    "case": ("infographics/case-study-template.png", (0, 0.03, 0, 0)),  # keep the "pilot data pending" stamp
    "tiers": ("infographics/pricing-tiers.png", (0, 0.2, 0, 0)),
    "founding": ("infographics/founding-member-offer.png", (0, 0.2, 0, 0)),
    "roadmap": ("infographics/roadmap-2026-2027.png", TOP),
    "regional": ("infographics/pricing-regional.png", (0, 0.2, 0, 0)),
    "matrix": ("infographics/payment-rails-matrix.png", (0, 0.19, 0, 0.06)),
}


def img(slide, key, alt, box_=(MX, 1.52, CW, 5.4), align="c", rounded=False):
    rel, crop = IMG[key]
    return pic(slide, rel, *box_, crop=crop, alt=alt, align=align, rounded=rounded)


# ================================================================ SLIDES
# Each builder: (prs, v, **kw) -> slide. v in {"master", "solo", "sme", "partner"}.

def s_title(prs, v):
    s = new_slide(prs, dark=True)
    pic(s, "brand/fluxmuse-logo-dark.png", MX, 0.55, 2.6, 0.9, align="tl", alt="FluxMuse logo")
    g = gid()
    box(s, 8.55, 1.25, 4.3, 4.3, fill=NIGHT_CARD, shape="oval", role="deco", group=g)
    pic(s, "brand/fluxmuse-icon-512.png", 8.95, 1.65, 3.5, 3.5, alt="FluxMuse Muse icon", group=g)
    if v == "partner":
        eyebrow, l1, l2 = "FluxMuse Partner Programme", "Serve more clients.", "Under your brand."
        sub = "White-label AI marketing and WhatsApp commerce for agencies and resellers."
    else:
        eyebrow, l1, l2 = "AI marketing & WhatsApp commerce", "Your AI marketing team.", "Selling inside WhatsApp."
        sub = "Plan, create, publish and get paid in chat. Built for South African businesses."
    text(s, MX, 2.05, 7.6, 0.3, eyebrow.upper(), size=13, bold=True, color=ORANGE, role="eyebrow", spc=150)
    text(s, MX, 2.42, 7.7, 1.9, [[(l1, {"color": PAPER})], [(l2, {"color": ORANGE})]], size=38, font=HEAD,
         bold=True, role="title", lsp=0.95)
    text(s, MX, 4.4, 7.2, 0.8, sub, size=18, color="D5DADF", role="body")
    g = gid()
    text(s, MX, 5.45, 7.4, 0.42, [[("Prepared for ", {"color": MUTED_DARK, "bold": False}),
                                   ("[[PROSPECT NAME]]", {"color": PAPER, "bold": True})]],
         size=18, role="placeholder", group=g)
    text(s, MX, 5.92, 7.4, 0.36, "[[DATE]]  ·  [[PRESENTER]]", size=15, color=MUTED_DARK, role="placeholder", group=g)
    notes(s, {
        "default": """Open warmly and make it about them. "Thanks for your time. FluxMuse gives a small business a
        full AI marketing team, and lets you sell and get paid right inside WhatsApp, where your customers already
        chat." Introduce yourself and say the session takes about 20 minutes plus questions. Ask one opening
        question: "Where do most of your sales conversations happen today?" Their answer tells you which slides to
        lean on. Swap guide: for owner-operators use the Solo deck, for businesses with 5 to 200 staff use the SME
        deck, and for agencies use the Partner Programme deck.""",
        "solo": """Keep it personal and short. "You're running the business, the marketing and the WhatsApp inbox
        yourself. FluxMuse gives you an AI marketing team from R499 a month, and lets you sell and get paid right
        inside WhatsApp." Introduce yourself and promise about 15 minutes. Ask: "How do customers find you, and how
        do they book or buy today?" Listen for lost DMs, late replies, no-shows and proof-of-payment screenshots;
        you'll come back to each one. If they have more than five staff or several brands, switch to the SME
        deck.""",
        "sme": """Frame this as a conversation about growth and control. "FluxMuse replaces a patchwork of marketing
        tools with one AI marketing team, and lets your team sell and get paid inside WhatsApp, with reporting that
        shows what drives revenue." Confirm who's in the room: owner, marketing lead, operations. Ask: "Which tools
        do you use today for social, WhatsApp, email and payments, and who runs each one?" Note the answers,
        because they set up the integrations slide. If they manage brands for clients, they're an agency: use the
        Partner deck.""",
        "partner": """This deck is for agencies, freelancers and resellers. "FluxMuse lets you offer AI marketing
        and WhatsApp commerce to your clients under your own brand, on a wholesale plan, and keep the spread."
        Introduce yourself and the partner programme. Ask two questions early: "How many clients do you manage?"
        and "What do you charge them per month?" Those numbers drive the illustrative economics later. Be clear
        from the start that partner obligations and referral commission are still being finalised, and don't
        promise terms that aren't on the slides.""",
    }, v)
    return s


def s_reality(prs, v):
    s = new_slide(prs)
    header(s, "The problem we solve", "The reality for South African businesses")
    img(s, "problem", "Today: disconnected tools, unaffordable retainers, payment-link drop-off and no local "
        "payment options. With FluxMuse: one AI team, plans from R499, checkout inside WhatsApp, 5 local rails.")
    notes(s, {
        "default": """Describe the day-to-day reality without numbers we can't back up. Most small businesses
        juggle separate tools for posts, chats, email, SMS and payments. Agency retainers are out of reach for
        many. Buyers drop off when they're sent to an external payment link, and checkout often doesn't offer
        local ways to pay like instant EFT or mobile money. On the right is what changes with FluxMuse: one AI
        marketing team, plans from R499 a month, checkout inside WhatsApp and five local payment rails. Ask which
        of the four problems hurts most, and focus the rest of the pitch there.""",
        "solo": """Talk about their week, not the market. You post when you find time, answer DMs late at night,
        and some buyers disappear when you send them a payment link. Paying a marketer or an agency isn't realistic
        yet, and plenty of buyers would rather pay by EFT than card, so you end up chasing proof-of-payment
        screenshots. On the right: FluxMuse gives you an AI marketing team from R499 a month, a catalog and
        checkout inside WhatsApp, and local ways to pay. Ask which of these costs them the most sales today.""",
        "sme": """For a business with staff, the problem is usually fragmentation. Posts, chats, email, SMS and
        payments sit in separate tools, often run by different people. Agency retainers are hard to justify
        without clear results, staff answer WhatsApp by hand, and nobody can connect social activity to sales.
        Buyers also drop off when sent to external payment links. On the right is the FluxMuse answer: one AI
        marketing team, checkout inside WhatsApp and five local payment rails. Ask how many tools they pay for
        today, and who owns each one.""",
    }, v)
    return s


def s_meet(prs, v):
    s = new_slide(prs, dark=True)
    g = gid("hdr")
    text(s, MX, 1.35, 4.3, 0.3, "MEET FLUXMUSE", size=12, bold=True, color=ORANGE, role="eyebrow", group=g, spc=150)
    text(s, MX, 1.72, 4.3, 2.3, "One AI team for marketing and sales", size=34, font=HEAD, bold=True,
         color=PAPER, role="title", group=g, lsp=0.95)
    text(s, MX, 4.1, 4.1, 1.0, "An AI marketing team and WhatsApp commerce platform, built for African SMBs.",
         size=17, color="D5DADF", role="body")
    g = gid()
    box(s, MX, 5.35, 3.95, 0.52, fill=ORANGE, shape="round", radius=0.26, role="chip", group=g)
    text(s, MX, 5.35, 3.95, 0.52, "14-day free trial · no card required", size=14, bold=True, color=NIGHT,
         align="c", anchor="m", role="chip", group=g)
    img(s, "hero", "FluxMuse website on a laptop and an order confirmed inside a WhatsApp chat on a phone",
        box_=(5.05, 0.75, 7.85, 5.95))
    notes(s, """Here's FluxMuse in one line: an AI marketing team and WhatsApp commerce platform built for African
    SMBs. On the laptop is our website; on the phone is the demo store confirming an order inside a WhatsApp chat.
    Everything in this deck is live in the product today unless it's labelled "in review" or "coming soon", and
    we'll be precise about which is which. There's a 14-day free trial with no card required, so trying it costs
    nothing. Pause here and ask whether they already use WhatsApp Business, and on which number.""", v)
    return s


def s_how(prs, v):
    s = new_slide(prs)
    header(s, "How FluxMuse works", "How it works: one loop, run by AI")
    img(s, "how", "Loop: Plan (Strategist), Create (Creator), Publish (Publisher), Sell (WhatsApp commerce), "
        "Learn (Analyst), with insights feeding the next plan")
    notes(s, """FluxMuse works as a loop run by four AI agents. The Strategist plans goals, channel mix and budget.
    The Creator writes copy and makes images in isiZulu, Afrikaans, Pidgin, Swahili, English and more. The
    Publisher schedules content across social, WhatsApp, email, SMS and USSD. Buyers then chat, browse and pay
    inside WhatsApp. The Analyst learns from orders and data and feeds insights back into the next plan. The point
    is that each step hands off to the next, so marketing and sales stop living in separate tools, and campaigns
    can go through approvals before anything is published.""", v)
    return s


def s_team(prs, v):
    s = new_slide(prs)
    header(s, "AI workforce", "Your AI marketing team")
    img(s, "roster", "Four core agents: Strategist, Creator, Publisher, Analyst; plus 23 specialist Flux agents "
        "and an AI Agent Marketplace")
    notes(s, """These are the four core agents: Strategist, Creator, Publisher and Analyst. Walk through one example
    for this prospect: the Strategist sets a goal for the month, the Creator drafts posts, the Publisher schedules
    them and runs the WhatsApp catalog, and the Analyst reports what worked. Behind them sit 23 specialist Flux
    agents for jobs like ads, nurture, support, commerce, finance and compliance, plus an AI Agent Marketplace.
    Don't read the chips aloud. Pick the two or three that match the prospect's pain, for example Flux Support for
    a busy inbox or Flux Nurture for repeat customers.""", v)
    return s


def s_sell(prs, v):
    s = new_slide(prs)
    header(s, "WhatsApp commerce", "Sell inside the chat")
    img(s, "flow", "Discover, Chat, Catalog, Cart, Pay with Yoco, Ozow, Paystack, pawaPay or Fincra, "
        "Confirmed, Loyalty", box_=(MX, 1.5, CW, 2.6))
    img(s, "phones", "Three WhatsApp screens: browse the catalog, review the cart, pay and get confirmation",
        box_=(MX, 4.25, 5.2, 2.68), align="l")
    rows = [("refresh", "Abandoned carts get an automatic recovery message."),
            ("card", "Buyers pay locally, with no redirect to a payment page."),
            ("users", "Segments and retargeting start the next sale.")]
    for i, (ic, t) in enumerate(rows):
        yy = 4.4 + i * 0.72
        g = icon(s, ic, "deep", 6.35, yy + 0.27, 0.54, circle=TINT)
        text(s, 6.8, yy, 5.93, 0.54, t, size=16, anchor="m", role="body", group=g)
    text(s, 6.1, 6.62, 6.6, 0.26, "All five rails live September 2026 (pawaPay and Fincra from 14 Sept).",
         size=11, color=MUTED, role="caption")
    notes(s, """This is the heart of the pitch. A buyer taps an ad, a post, a QR code or a wa.me link and lands in a
    chat. A chatbot or WhatsApp Flow replies instantly, they browse the catalog, build a cart and pay without
    leaving WhatsApp, using Yoco, Ozow, Paystack, pawaPay or Fincra. The order is confirmed and status updates
    arrive in the same chat. If they abandon the cart, a recovery message goes out automatically, and loyalty and
    retargeting bring them back. If you demo live, stick to the catalog, cart and confirmation screens shown
    here.""", v)
    return s


def s_channels(prs, v):
    s = new_slide(prs)
    if v == "sme":
        header(s, "Platform", "One platform, from channel to checkout")
        img(s, "stack", "Six layers: channels (Instagram in Meta review), AI workforce, commerce and payments, "
            "growth and CRM, integrations, trust")
    else:
        header(s, "Channels", "One platform, every channel", w=7.6)
        text(s, 8.35, 0.74, 4.38, 0.62, "Instagram: in Meta review. TikTok, LinkedIn, X, YouTube, Threads and "
             "Pinterest: coming soon.", size=11, color=MUTED, role="caption", align="r")
        img(s, "hub", "FluxMuse AI team connected to WhatsApp, Facebook Pages, Messenger, email, SMS, USSD, website "
            "chat widget, landing pages and link-in-bio; Instagram in Meta review")
    notes(s, {
        "default": """One AI team runs every channel your buyers use. Live today: WhatsApp, Facebook Pages,
        Messenger, email, SMS, USSD, a website chat widget, landing pages and link-in-bio. Be precise about
        Instagram: it's in Meta App Review, so publishing, comments and DMs aren't live yet. TikTok, LinkedIn, X,
        YouTube, Threads and Pinterest are coming soon as part of our social adapters work, with no dates promised.
        If Instagram is their main channel, say so honestly and start them on WhatsApp and Facebook, adding
        Instagram once Meta approves it.""",
        "sme": """This is the whole platform in six layers. Channels: WhatsApp, Facebook Pages, Messenger, email,
        SMS, USSD, web chat and link-in-bio, with Instagram in Meta review and other networks coming soon. The AI
        workforce: four core agents plus 23 specialists. Commerce and payments: catalog, checkout in chat, orders
        and abandoned-cart recovery on five rails. Growth and CRM: lead scoring, deals, segments, loyalty and
        churn-risk scoring. Integrations with Shopify, WooCommerce, Takealot, HubSpot, accounting, Slack and
        Zapier. Trust controls sit underneath. Ask which layer replaces the most tools for them today.""",
    }, v)
    return s


def s_paid(prs, v):
    s = new_slide(prs)
    header(s, "Payments coverage", "Get paid locally")
    img(s, "map", "Tile map of 23 African countries covered by 5 live payment rails, with Botswana and Namibia "
        "coming soon")
    notes(s, """Getting paid is where many chat sales die, so this matters. FluxMuse has five payment rails: Yoco
    for cards and tap-to-pay, Ozow for instant EFT, Paystack, pawaPay for mobile money, and Fincra for collections
    and payouts. Together they cover 23 African countries, all live from September 2026, with pawaPay and Fincra
    from 14 September. Botswana and Namibia are coming soon, once a rail covers them. For a South African business
    the headline is simple: card, tap-to-pay and instant EFT, inside the chat. Save the country detail for the
    appendix.""", v)
    return s


def s_segments(prs, v):
    s = new_slide(prs, dark=True)
    g = gid("hdr")
    text(s, MX, 0.85, CW, 0.3, "BUILT FOR YOU", size=12, bold=True, color=ORANGE, role="eyebrow", group=g, spc=150)
    text(s, MX, 1.15, CW, 0.8, "Built for businesses like yours", size=36, font=HEAD, bold=True, color=PAPER,
         role="title", group=g)
    text(s, MX, 1.98, CW, 0.4, "South Africa first.", size=18, color=MUTED_DARK, role="body")
    cards = [("users", "Solo entrepreneurs", "Founder-run, 1–5 people", "Starter · R499/mo"),
             ("briefcase", "SMEs", "5–200 staff", "Growth · R1,999/mo"),
             ("layers", "Agencies & partners", "Managing 5–50 clients", "Wholesale · R5,599/mo")]
    cwid, gap = 3.84, (CW - 3 * 3.84) / 2
    for i, (ic, name, who, tier) in enumerate(cards):
        x = MX + i * (cwid + gap)
        g = gid("seg")
        box(s, x, 2.85, cwid, 3.7, fill=NIGHT_CARD, shape="round", radius=0.2, role="card", group=g)
        icon(s, ic, "ink", x + 0.8, 3.65, 0.9, circle=ORANGE, group=g)
        # name box is bottom-anchored with room for two lines, so wrapped names grow upwards
        text(s, x + 0.35, 4.25, cwid - 0.7, 0.9, name, size=22, font=HEAD, bold=True, color=PAPER, role="card",
             group=g, anchor="b", lsp=0.95)
        text(s, x + 0.35, 5.25, cwid - 0.7, 0.4, who, size=16, color=MUTED_DARK, role="card", group=g)
        text(s, x + 0.35, 5.8, cwid - 0.7, 0.45, tier, size=18, bold=True, color=ORANGE, role="card", group=g)
    notes(s, """FluxMuse is built for three kinds of business, all starting in South Africa: the Gauteng pilot first,
    then national, before Nigeria, Kenya and Ghana. Solo entrepreneurs run founder-led businesses of one to five
    people and usually start on Starter at R499 a month. SMEs have five to 200 staff and typically run Growth at
    R1,999, moving up to Scale. Agencies and freelancers managing five to 50 clients join the partner programme on
    the Agency tier at a wholesale R5,599 a month. Ask which one fits them, then go to that segment's slide.""", v)
    return s


def s_seg_solo(prs, v):
    s = new_slide(prs)
    header(s, "Launch segment · Solo entrepreneurs", "Solo entrepreneurs: 1–5 people")
    img(s, "solo", "Solo entrepreneurs: pains, FluxMuse answers, Starter plan price and Founding Member price")
    notes(s, """Solo entrepreneurs are side hustles, braiders and hair studios, fashion resellers, home bakers,
    coaches and informal retailers selling on Instagram, Facebook and WhatsApp. Their pains are simple: no time or
    budget for a marketer, sales lost in busy DMs, and buyers who drop off at payment links. FluxMuse answers each
    one: the Creator writes posts and the Publisher schedules them, chatbots reply instantly even after hours, and
    checkout happens in the chat. Starter is R499 a month; the full catalog and checkout sit in Growth, the natural
    upgrade. Founding Members pay R349 for their first two Starter bills.""", v)
    return s


def s_seg_sme(prs, v):
    s = new_slide(prs)
    header(s, "Launch segment · SMEs", "SMEs: 5–200 staff")
    img(s, "sme", "SMEs: pains, FluxMuse answers, Growth plan price and Founding Member price")
    notes(s, """SMEs are retailers, restaurants and franchises, e-commerce brands, clinics, property, auto dealers,
    education and professional services, with five to 200 staff. The usual story: six or more disconnected tools,
    costly agency retainers, no clear link from social activity to sales, and staff answering WhatsApp by hand.
    FluxMuse replaces the stack, lets the team sell and get paid in chat, and shows what drives revenue. Growth at
    R1,999 a month is the usual starting point, with Scale at R4,999 for multi-brand or multi-location businesses.
    Founding Members pay R1,399 for their first two Growth bills.""", v)
    return s


def s_seg_agency(prs, v):
    s = new_slide(prs)
    header(s, "Launch segment · Agencies & partners", "Agencies: serve more clients, your brand")
    img(s, "agency", "Agencies: pains, FluxMuse answers, Agency tier and partner wholesale price")
    notes(s, """Agencies, digital and social media shops and freelancers managing five to 50 SMB clients feel a margin
    squeeze, spend hours on manual reporting, juggle too many tools per client, and have clients asking for
    WhatsApp commerce and AI. FluxMuse gives them one white-label platform with 50 client sub-accounts and reseller
    billing. Partners buy the Agency tier at a permanent wholesale price of R5,599 a month, 30% off the R7,999 list
    price, and set their own retail price. Don't quote a fixed margin percentage: margin depends on the agency's
    pricing and delivery costs.""", v)
    return s


def s_partner_model(prs, v):
    s = new_slide(prs)
    header(s, "Agency partner model · illustrative economics", "Resell FluxMuse under your own brand")
    img(s, "partner", "FluxMuse to agency on the wholesale Agency tier to clients at a retail price the agency sets; "
        "illustrative example of 20 clients")
    notes(s, """Here's how the partner model works. FluxMuse invoices the agency at the wholesale Agency-tier price of
    R5,599 a month. The agency serves its clients under its own brand and bills them a retail price it sets. The
    example is illustrative only, not a forecast: 20 clients at R1,500 a month bills R30,000, and after R5,599 the
    agency keeps a gross spread of R24,401 a month, before its own service costs. Encourage them to run the numbers
    with their real client count and prices. Referral commission for clients on their own plans is still to be
    decided.""", v)
    return s


def s_action(prs, v):
    s = new_slide(prs)
    header(s, "Product", "See it in action")
    img(s, "laptop", "FluxMuse website home page on a laptop", box_=(0.45, 1.75, 6.95, 4.2), align="t")
    text(s, 0.6, 6.05, 6.5, 0.28, "fluxmuse.ai home page", size=11, color=MUTED, role="caption")
    img(s, "pricing_shot", "FluxMuse pricing page in rand: Starter R499, Growth R1,999, Scale R4,999, Agency R7,999",
        box_=(7.7, 1.62, 5.03, 1.0), align="t", rounded=True)
    text(s, 7.7, 2.66, 5.03, 0.28, "Pricing page, in rand", size=11, color=MUTED, role="caption")
    labels = [("ph1", "1 · Catalog"), ("ph3", "2 · Cart"), ("ph7", "3 · Paid")]
    pw, gap = 1.53, (5.03 - 3 * 1.53) / 2
    for i, (k, lab) in enumerate(labels):
        x = 7.7 + i * (pw + gap)
        g = gid()
        img(s, k, f"WhatsApp checkout demo, step {lab}", box_=(x, 3.1, pw, 3.12), align="t")
        text(s, x, 6.3, pw, 0.28, lab, size=11, color=MUTED, align="c", role="caption", group=g)
    notes(s, """Show, don't tell. On the left is the FluxMuse website; top right is the pricing page in rand; below it
    are three screens from the WhatsApp checkout demo: browse the catalog, review the cart, and pay with a
    confirmation in the chat. If there's time and a good connection, open the live demo and walk the same three
    steps, but skip the payment-method screens, which still list a provider outside our five rails. Invite them to
    start a free trial during the call so they can click through their own workspace while you talk.""", v)
    return s


def s_journey(prs, v):
    s = new_slide(prs)
    header(s, "Working with FluxMuse", "How we work together")
    img(s, "journey", "Discovery call, audit and strategy, setup, pilot campaign, report and optimise, convert "
        "to a plan; KPIs tracked from day one")
    notes(s, """This is how an engagement runs. Week 0 is a discovery call on goals, audience, products and channels.
    In week 1 we audit channels and content, agree a pilot plan and set KPI goals. Setup takes one to two weeks:
    WhatsApp Business number, catalog, payment rail and content calendar. Then a pilot campaign, typically 60 days
    and up to 90, with weekly check-ins. We report against the goals monthly, and when the pilot ends you move onto
    a plan. Stress that KPI targets are goals we agree together, not promised results.""", v)
    return s


def s_pilot(prs, v):
    s = new_slide(prs)
    header(s, "Gauteng pilot · underway · results December 2026", "Proof: the Gauteng pilot")
    img(s, "pilot", "Gauteng pilot: 12 brands, September to 30 November 2026, suggested segment mix, KPIs tracked, "
        "results December 2026")
    notes(s, """We're proving this in the market right now. Twelve brands in Gauteng are running FluxMuse from
    September to 30 November 2026, with a suggested mix of solo entrepreneurs, SMEs and agencies. We track brands
    onboarded, WhatsApp conversations, orders and GMV, content published, time saved, conversion uplift and NPS. Be
    honest: the pilot is underway and results arrive in December 2026. We don't share projected results, and we
    only publish figures with each brand's written consent. Offer to send them the results pack in December.""", v)
    return s


def s_case(prs, v):
    s = new_slide(prs)
    header(s, "Lead case study · pilot data pending", "A case study in the making")
    img(s, "case", "Case study template for a Gauteng braiding and hair studio; every result is a placeholder "
        "until pilot data is available")
    notes(s, """Our lead case study is a Gauteng braiding and hair studio: an owner-operator on Starter who already
    books clients on WhatsApp, with short booking cycles, so 60 days is enough to see a difference. In week one we
    set up a booking chatbot, a catalog with a deposit link through Yoco or Ozow, appointment reminders,
    multilingual before-and-after content, and a wa.me QR code for the salon mirror. Every result on this slide is a
    placeholder until the pilot ends, and the targets are goals, not outcomes. Replace this slide with the consented
    case study from December 2026.""", v)
    return s


def s_pricing(prs, v):
    s = new_slide(prs)
    header(s, "Pricing (ZAR) · 14-day free trial, no card", "Simple plans, priced in rand")
    img(s, "tiers", "Four plans in rand: Starter R499, Growth R1,999, Scale R4,999, Agency R7,999 per month, "
        "annual 10 times monthly")
    notes(s, {
        "default": """Four plans, priced in rand, all with a 14-day free trial and no card required. Starter at R499
        a month suits solo founders: one brand and three channels. Growth at R1,999 is the most popular: three
        brands, 15 channels, e-commerce and WhatsApp commerce, USSD and integrations. Scale at R4,999 adds API
        access, white-label reports and priority support for ten brands. Agency at R7,999 brings full white-label,
        50 sub-accounts and reseller billing. Annual billing is ten times monthly, so two months free. Enterprise is
        available on request.""",
        "solo": """Most solo businesses start on Starter: R499 a month for one brand, three channels and 5,000 AI
        credits, with community support. When they're ready for e-commerce and full WhatsApp commerce with a
        catalog and checkout, Growth at R1,999 is the step up. Point out the 14-day free trial with no card
        required, and annual billing at ten times monthly, which gives two months free. Scale and Agency are for
        bigger teams, so don't dwell on them. Then move to the Founding Member offer, which makes the first two
        months cheaper.""",
        "sme": """Most SMEs start on Growth at R1,999 a month: three brands, 15 channels, 25,000 AI credits,
        e-commerce and WhatsApp commerce, USSD, integrations and email support. Multi-brand or multi-location
        businesses should look at Scale at R4,999: ten brands, 40 channels, API access, white-label reports and
        priority support. Every plan starts with a 14-day free trial, no card required, and annual billing is ten
        times monthly, so two months free. Larger groups can ask about Enterprise, but treat those conversations as
        inbound.""",
    }, v)
    return s


def s_pricing_table(prs, v):
    s = new_slide(prs)
    header(s, "Pricing (ZAR) · editable", "Plans at a glance")
    hdr = {"fill": NIGHT, "color": PAPER, "bold": True, "font": HEAD, "size": 16, "line": None}
    grow = lambda t, **k: {"text": t, "fill": TINT, "bold": k.get("bold", False)}
    lab = lambda t: {"text": t, "bold": True, "align": "l"}
    rows = [
        [dict(hdr, text="Plan", align="l"), dict(hdr, text="Starter"), dict(hdr, text="Growth · most popular", fill=ORANGE, color=NIGHT),
         dict(hdr, text="Scale"), dict(hdr, text="Agency")],
        [lab("Monthly"), {"text": "R499", "bold": True}, grow("R1,999", bold=True), {"text": "R4,999", "bold": True}, {"text": "R7,999", "bold": True}],
        [lab("Annual (10× monthly)"), "R4,990", grow("R19,990"), "R49,990", "R79,990"],
        [lab("Founding Member, first 2 bills"), "R349", grow("R1,399"), "R3,499", "Partner pricing"],
        [lab("Brands · channels"), "1 · 3", grow("3 · 15"), "10 · 40", "Unlimited · 80"],
        [lab("AI credits / month"), "5,000", grow("25,000"), "100,000", "500,000"],
        [lab("Best for"), "Solo founders", grow("E-commerce & WhatsApp"), "API & white-label reports", "White-label, 50 sub-accounts"],
    ]
    for c in rows[0]:
        c["size"] = 15
    table(s, MX, 1.62, [2.93, 2.3, 2.3, 2.3, 2.3], [0.74, 0.52, 0.52, 0.74, 0.52, 0.52, 0.74], rows, size=15,
          label="pricing-zar")
    g = gid()
    box(s, MX, 6.12, 4.2, 0.5, fill=ORANGE, shape="round", radius=0.25, role="chip", group=g)
    text(s, MX, 6.12, 4.2, 0.5, "14-day free trial · no card required", size=14, bold=True, color=NIGHT, align="c",
         anchor="m", role="chip", group=g)
    text(s, 5.1, 6.14, 7.63, 0.5, "Enterprise from R19,999/mo: talk to us. Nigeria, Kenya, Ghana and USD pricing: "
         "see appendix.", size=12, color=MUTED, anchor="m", role="caption")
    notes(s, """This is the same pricing as an editable table, useful when you tailor the deck. It shows monthly and
    annual prices, with annual at ten times monthly, and what each plan includes. The Founding Member row is the
    launch price for the first two monthly bills in the South African launch window: R349, R1,399 and R3,499.
    Agencies aren't offered Founding Member pricing; point them to partner wholesale pricing instead. If you edit
    any price, check it against the facts file first. Enterprise starts from R19,999 a month, and those
    conversations are inbound only.""", v)
    return s


def s_founding(prs, v):
    s = new_slide(prs)
    header(s, "Launch offer · South Africa · 1 Dec 2026 – 31 Jan 2027", "Founding Member launch offer")
    img(s, "founding", "Founding Member offer: 30% off the first 2 monthly bills or 2 extra months on annual plans, "
        "badge and priority support, South Africa launch window 1 Dec 2026 to 31 Jan 2027",
        box_=(MX, 1.52, CW, 4.65))
    chips = [("Starter R349", TINT, INK), ("Growth R1,399", TINT, INK), ("Scale R3,499", TINT, INK),
             ("Annual: +2 months free", MIST, INK), ("Agencies: partner pricing", MIST, INK)]
    widths = [2.0, 2.2, 2.0, 2.85, 2.9]
    gap = (CW - sum(widths)) / 4
    x = MX
    for (t, f, c), wdt in zip(chips, widths):
        g = gid()
        box(s, x, 6.3, wdt, 0.52, fill=f, shape="round", radius=0.26, role="chip", group=g)
        text(s, x, 6.3, wdt, 0.52, t, size=15, bold=True, color=c, align="c", anchor="m", role="chip", group=g)
        x += wdt + gap
    notes(s, {
        "default": """Our launch offer is Founding Member. Sign up in the South African launch window, 1 December
        2026 to 31 January 2027, and get 30% off your first two monthly bills: R349 on Starter, R1,399 on Growth or
        R3,499 on Scale, then list price. On an annual plan you get two extra months free. Founding Members also get
        a Founding Member badge and priority support. It comes on top of the 14-day free trial. Agencies aren't part
        of this offer; they get partner wholesale pricing instead. Don't promise how long the perks last.""",
        "solo": """Here's the easiest yes. Start the 14-day free trial now, then subscribe in the launch window, 1
        December 2026 to 31 January 2027, and your first two Starter bills are R349 instead of R499. If you move to
        Growth for the catalog and checkout, it's R1,399 for two months, then R1,999. Pay annually and you get two
        extra months free. Founding Members also get a badge and priority support. Don't promise how long the perks
        last, and don't offer the deeper pilot pricing, which is only for the twelve Gauteng pilot brands.""",
    }, v)
    return s


def s_trust(prs, v):
    s = new_slide(prs)
    header(s, "Trust & compliance", "Built to protect your customers' data")
    cards = [("shield", "POPIA, NDPR, GDPR", "Ready: South Africa, Nigeria, EU"),
             ("checkCircle", "Meta Tech Provider", "Verified by Meta"),
             ("layers", "Row-level security", "Every database table"),
             ("plug", "Encrypted tokens", "Connected accounts"),
             ("file", "Audit export", "Exportable activity record"),
             ("refresh", "Data export & deletion", "Export data, delete account")]
    cwid, ch = 3.84, 2.42
    gapx = (CW - 3 * cwid) / 2
    for i, (ic, t, sub) in enumerate(cards):
        x = MX + (i % 3) * (cwid + gapx)
        y = 1.68 + (i // 3) * (ch + 0.22)
        g = gid("trust")
        box(s, x, y, cwid, ch, fill=MIST, shape="round", radius=0.18, role="card", group=g)
        icon(s, ic, "deep", x + 0.68, y + 0.68, 0.78, circle=PAPER, group=g)
        text(s, x + 0.32, y + 1.2, cwid - 0.6, 0.5, t, size=18, font=HEAD, bold=True, role="card", group=g)
        text(s, x + 0.32, y + 1.7, cwid - 0.6, 0.45, sub, size=14, color=MUTED, role="card", group=g)
    notes(s, """Trust matters when you're handling customer chats and payments. FluxMuse is POPIA, NDPR and GDPR ready.
    Fluxmuse Pty Ltd has passed Meta Business Verification and is a verified Meta Tech Provider, and WhatsApp runs on
    Meta's official Cloud API. Under the hood there's row-level security on every database table, encrypted tokens
    for connected accounts, audit export, and account deletion and data export. If they ask for a data processing
    agreement, a security questionnaire or certifications, don't improvise: note the request and follow up in
    writing.""", v)
    return s


def s_roadmap(prs, v):
    s = new_slide(prs)
    header(s, "Roadmap · indicative", "What's coming")
    img(s, "roadmap", "Indicative roadmap 2026 to 2028: markets and payments lane and product lane")
    notes(s, {
        "default": """Here's what's coming — indicative, not a commitment. On markets: pawaPay and Fincra go
        live in September 2026, the Gauteng pilot runs to 30 November, and South Africa launches nationally with the
        Founding Member offer from 1 December. 2027 scales South Africa, with no new launches. Nigeria
        follows in February 2028, Kenya in April 2028 and Ghana in June 2028, then the rest of the rail-covered
        countries in US dollars from November 2028. Botswana and Namibia are coming soon. Those launch
        months are milestone-gated in the financial model, so they move with revenue. On product, in order: Meta
        permissions including Instagram, targeted for Q4 2026, then social adapters, business app integrations,
        then the Flux_Partner programme. Don't give dates beyond these.""",
        "partner": """Here's what's coming, and it's indicative, not a commitment. On markets: pawaPay and Fincra go
        live in September 2026, the Gauteng pilot runs to 30 November, and South Africa launches nationally from 1
        December. 2027 scales South Africa; Nigeria follows in February 2028, Kenya in April 2028, Ghana in June
        2028, then the rest of the rail-covered countries in US dollars from November 2028. Those months are
        milestone-gated in the financial model, so they move with revenue. On product, in order: Meta permissions
        including Instagram, targeted for Q4 2026, then social adapters, then business app integrations, then the
        Flux_Partner programme with referral tracking and partner-managed workspaces. White-label, sub-accounts and
        reseller billing are already live today; the programme layer comes later.""",
    }, v)
    return s


def s_next(prs, v):
    s = new_slide(prs, dark=True)
    g = gid("hdr")
    text(s, MX, 0.8, CW, 0.3, "NEXT STEPS", size=12, bold=True, color=ORANGE, role="eyebrow", group=g, spc=150)
    title = "Let's build your partner practice" if v == "partner" else "Let's get you selling in chat"
    text(s, MX, 1.1, CW, 0.8, title, size=36, font=HEAD, bold=True, color=PAPER, role="title", group=g)
    ctas = {
        "default": [("Book a 30-minute discovery call", "We map your goals, channels and products."),
                    ("Start your 14-day free trial", "No card required, at fluxmuse.ai."),
                    ("Become a partner", "Agencies: Agency tier at R5,599/mo wholesale.")],
        "solo": [("Book a 30-minute discovery call", "We map your goals, channels and products."),
                 ("Start your 14-day free trial", "No card required, at fluxmuse.ai."),
                 ("Join as a Founding Member", "First 2 Starter bills at R349.")],
        "sme": [("Book a 30-minute discovery call", "Goals, channels, tools and team."),
                ("Start your 14-day free trial", "No card required, at fluxmuse.ai."),
                ("Get a tailored proposal", "Multi-brand, Scale or a campaign pilot.")],
        "partner": [("Book a 30-minute partner call", "Your clients, services and pricing."),
                    ("Start your 14-day free trial", "No card required, at fluxmuse.ai."),
                    ("Apply to become a partner", "Agency tier at R5,599/mo wholesale.")],
    }
    for i, (t, sub) in enumerate(ctas.get(v, ctas["default"])):
        y = 2.25 + i * 1.42
        g = gid("cta")
        box(s, MX, y, 7.55, 1.22, fill=NIGHT_CARD, shape="round", radius=0.18, role="card", group=g)
        box(s, MX + 0.32, y + 0.3, 0.62, 0.62, fill=ORANGE, shape="oval", role="deco", group=g)
        text(s, MX + 0.32, y + 0.3, 0.62, 0.62, str(i + 1), size=20, font=HEAD, bold=True, color=NIGHT, align="c",
             anchor="m", role="card", group=g)
        text(s, MX + 1.2, y + 0.2, 6.15, 0.45, t, size=20, font=HEAD, bold=True, color=PAPER, role="card", group=g)
        text(s, MX + 1.2, y + 0.7, 6.15, 0.38, sub, size=15, color=MUTED_DARK, role="card", group=g)
    g = gid("contact")
    box(s, 8.5, 2.25, 4.23, 4.06, fill=NIGHT_CARD, shape="round", radius=0.18, role="card", group=g)
    text(s, 8.85, 2.5, 3.6, 0.3, "TALK TO US", size=12, bold=True, color=ORANGE, role="eyebrow", group=g, spc=150)
    text(s, 8.85, 2.85, 3.6, 1.3, [{"runs": "[[CONTACT NAME]]", "bold": True, "color": PAPER, "space_after": 4},
                                   {"runs": "[[EMAIL]]", "space_after": 4}, "[[WHATSAPP NUMBER]]"],
         size=16, color="D5DADF", role="placeholder", group=g)
    box(s, 8.85, 4.3, 1.6, 1.6, fill=None, line=MUTED_DARK, lw=1.25, dash=True, shape="round", radius=0.1,
        role="deco", group=g)
    text(s, 8.85, 4.3, 1.6, 1.6, "[[wa.me QR CODE]]", size=11, color=MUTED_DARK, align="c", anchor="m",
         role="caption", group=g)
    text(s, 10.62, 4.45, 1.95, 1.4, [{"runs": "Scan to chat on WhatsApp", "space_after": 6},
                                     {"runs": "wa.me/", "color": PAPER, "bold": True},
                                     {"runs": "[[NUMBER]]", "color": PAPER, "bold": True}],
         size=14, color=MUTED_DARK, role="placeholder", group=g)
    notes(s, {
        "default": """Close with one clear next step, and agree it before the call ends. Option one: book a 30-minute
        discovery call to map goals, channels and products. Option two: start the 14-day free trial today at
        fluxmuse.ai, with no card needed. Option three, for agencies: apply to become a partner and get the Agency
        tier at wholesale pricing. Share your contact details and the WhatsApp QR code so they can message you
        directly. Follow up within 24 hours with this deck and a short summary of what you agreed.""",
        "solo": """Close with one clear next step. The simplest is to start the 14-day free trial right now, on the
        call, at fluxmuse.ai; no card is needed, and you can help them connect WhatsApp and add their first
        products. If they want help first, book a 30-minute discovery call. Remind them that joining as a Founding
        Member in the launch window means R349 for their first two Starter bills. Share your WhatsApp QR code so they
        can message you with questions, and follow up within 24 hours.""",
        "sme": """Close with one agreed next step and an owner for it. Option one: a 30-minute discovery call with
        the people who run marketing, sales and operations, to map goals, channels, tools and team. Option two:
        start the 14-day free trial now, with no card, so the team can explore. Option three: we prepare a tailored
        proposal for multi-brand, Scale or a campaign pilot. Share your contact details and WhatsApp QR code, and
        follow up within 24 hours with the deck and a short written summary.""",
        "partner": """Close with a concrete partner next step. Option one: book a 30-minute partner call to go
        through their clients, services and pricing in detail. Option two: start a 14-day free trial so they can
        explore the platform before committing. Option three: apply to become a partner and move onto the Agency
        tier at R5,599 a month wholesale, starting with one pilot client. Share your contact details and WhatsApp QR
        code. Follow up within 24 hours, and don't commit to commission or obligations that are still to be
        confirmed.""",
    }, v)
    return s


# ---------------------------------------------------------------- SME-only
def s_sme_integrations(prs, v):
    s = new_slide(prs)
    header(s, "Integrations & teamwork", "Fits your stack. Works for your team.")
    g = gid("intg")
    box(s, MX, 1.68, 5.95, 5.05, fill=MIST, shape="round", radius=0.2, role="card", group=g)
    icon(s, "plug", "deep", MX + 0.62, 2.22, 0.66, circle=PAPER, group=g)
    text(s, MX + 1.1, 1.97, 4.5, 0.5, "Connects to", size=20, font=HEAD, bold=True, role="card", group=g)
    chips = ["Shopify", "WooCommerce", "Takealot", "HubSpot", "Accounting sync", "Slack", "Zapier", "Public API"]
    for i, c in enumerate(chips):
        x = MX + 0.35 + (i % 2) * 2.7
        y = 2.9 + (i // 2) * 0.8
        gg = gid()
        box(s, x, y, 2.5, 0.6, fill=PAPER, shape="round", radius=0.3, role="chip", group=gg)
        text(s, x, y, 2.5, 0.6, c, size=15, bold=True, align="c", anchor="m", role="chip", group=gg)
    text(s, MX + 0.35, 6.15, 5.3, 0.3, "Integrations from Growth · API access from Scale", size=12, color=MUTED,
         role="caption", group=g)
    g = gid("team")
    box(s, 6.78, 1.68, 5.95, 5.05, fill=TINT, shape="round", radius=0.2, role="card", group=g)
    icon(s, "users", "deep", 6.78 + 0.62, 2.22, 0.66, circle=PAPER, group=g)
    text(s, 6.78 + 1.1, 1.97, 4.5, 0.5, "Built for teams", size=20, font=HEAD, bold=True, role="card", group=g)
    rows = [("checkCircle", "Campaign approvals"), ("layers", "Up to 10 brands (Scale)"),
            ("trend", "CRM deals & lead scoring"), ("chart", "Custom reports, A/B tests"),
            ("chat", "Sentiment escalation")]
    for i, (ic, t) in enumerate(rows):
        y = 2.92 + i * 0.72
        gg = icon(s, ic, "deep", 7.4, y + 0.28, 0.5, circle=PAPER)
        text(s, 7.85, y, 4.7, 0.56, t, size=16, anchor="m", role="body", group=gg)
    notes(s, """SMEs rarely start from zero, so show how FluxMuse fits what they already run. On the left are the
    integrations: Shopify, WooCommerce and Takealot for commerce, HubSpot and CRM sync, accounting sync, Slack,
    Zapier, and a public API with keys. Integrations are included from Growth, and API access from Scale. On the
    right is what helps a team work together: campaigns with approvals, several brands on one account, CRM deals,
    lead scoring and segments, custom reports and A/B testing, and sentiment analysis that escalates unhappy
    customers. Ask which tools they'd want connected first.""", v)
    return s


# ---------------------------------------------------------------- partner-only
def p_squeeze(prs, v):
    s = new_slide(prs)
    header(s, "The agency squeeze", "What agencies are up against")
    today = [("x", "Margin squeeze on retainers"), ("x", "Manual reporting"), ("x", "Too many tools per client")]
    after = [("check", "Wholesale price, your retail price"), ("check", "White-label, 50 sub-accounts"),
             ("check", "WhatsApp commerce + AI, your brand")]
    cwid = 3.84
    gapx = (CW - 3 * cwid) / 2
    text(s, MX, 1.62, 4, 0.28, "TODAY", size=12, bold=True, color=MUTED, role="eyebrow", spc=150)
    text(s, MX, 4.12, 4, 0.28, "WITH FLUXMUSE", size=12, bold=True, color=DEEP, role="eyebrow", spc=150)
    for row, (items, fill, col, y) in enumerate([(today, MIST, "ink", 1.98), (after, TINT, "deep", 4.48)]):
        for i, (ic, t) in enumerate(items):
            x = MX + i * (cwid + gapx)
            g = gid("sq")
            box(s, x, y, cwid, 1.6, fill=fill, shape="round", radius=0.18, role="card", group=g)
            icon(s, ic, col, x + 0.62, y + 0.8, 0.66, circle=PAPER, group=g)
            text(s, x + 1.15, y + 0.2, cwid - 1.4, 1.2, t, size=18, font=HEAD, bold=True, anchor="m", role="card", group=g)
            if row == 0:
                box(s, x + cwid / 2 - 0.2, 3.66, 0.4, 0.38, fill=ORANGE, shape="down", role="deco", group=gid())
    notes(s, """Start with the agency's reality. Clients expect more every month, from content to WhatsApp sales to AI,
    while retainers stay flat. Reporting eats hours, and every new client adds another set of tools and logins.
    FluxMuse flips that. Partners buy the Agency tier at wholesale and set their own retail price. They run every
    client from one white-label platform with 50 client sub-accounts. And they resell AI agents and WhatsApp
    checkout under their own brand. Ask how many tools they currently pay for per client, and how long monthly
    reporting takes them.""", v)
    return s


def p_what_you_get(prs, v):
    s = new_slide(prs)
    header(s, "Partner plan · Agency tier at wholesale", "Everything you need to resell")
    g = gid("price")
    box(s, MX, 1.68, 4.05, 5.05, fill=TINT, shape="round", radius=0.2, role="card", group=g)
    text(s, MX + 0.4, 2.0, 3.3, 0.36, "PARTNER WHOLESALE", size=13, bold=True, color=DEEP, role="eyebrow", group=g, spc=120)
    text(s, MX + 0.4, 2.4, 3.4, 1.2, "R5,599", size=54, font=HEAD, bold=True, role="card", group=g)
    text(s, MX + 0.4, 3.65, 3.3, 0.7, "per month, billed to your agency", size=16, role="card", group=g)
    text(s, MX + 0.4, 4.5, 3.3, 0.7, "30% off the R7,999 Agency list price", size=16, bold=True, role="card", group=g)
    text(s, MX + 0.4, 5.4, 3.3, 0.9, "Annual: R55,990 (10× monthly). Launch offers don't stack.", size=14,
         color=MUTED, role="card", group=g)
    items = [("layers", "Full white-label"), ("users", "50 client sub-accounts"), ("wallet", "Reseller billing"),
             ("globe", "80 channels"), ("sparkles", "500,000 AI credits/mo"), ("chat", "Dedicated support")]
    x0, cwid, ch = 4.95, 3.74, 1.5
    for i, (ic, t) in enumerate(items):
        x = x0 + (i % 2) * (cwid + 0.3)
        y = 1.68 + (i // 2) * (ch + 0.275)
        g = gid("inc")
        box(s, x, y, cwid, ch, fill=PAPER, line=LINE, lw=1.25, shape="round", radius=0.18, role="card", group=g)
        icon(s, ic, "deep", x + 0.65, y + ch / 2, 0.72, circle=TINT, group=g)
        text(s, x + 1.2, y + 0.2, cwid - 1.4, ch - 0.4, t, size=18, font=HEAD, bold=True, anchor="m", role="card", group=g)
    notes(s, """This is what a partner gets on the Agency tier at wholesale pricing: R5,599 a month, 30% off the R7,999
    list price, or R55,990 a year. Full white-label, so clients see the agency's brand. Fifty client sub-accounts,
    one per client. Reseller billing, so the agency bills its clients directly. Eighty channels and 500,000 AI
    credits a month, plus dedicated support. Clients get the same product as direct customers: the AI marketing
    team, WhatsApp commerce and local payment rails. Launch offers don't stack with wholesale pricing.""", v)
    return s


def wholesale_table(s, y):
    hdr = {"fill": NIGHT, "color": PAPER, "bold": True, "font": HEAD, "size": 16, "line": None}
    ws = lambda t: {"text": t, "fill": TINT, "bold": True}
    rows = [
        [dict(hdr, text="", align="l"), dict(hdr, text="ZAR"), dict(hdr, text="NGN"), dict(hdr, text="KES"),
         dict(hdr, text="GHS"), dict(hdr, text="USD")],
        [{"text": "Agency list / month", "bold": True}, "R7,999", "₦662,000", "KSh 64,499", "GH₵ 5,439", "$429"],
        [{"text": "Wholesale / month (−30%)", "bold": True, "fill": TINT}, ws("R5,599"), ws("₦463,000"),
         ws("KSh 44,999"), ws("GH₵ 3,799"), ws("$299")],
        [{"text": "Wholesale / year", "bold": True}, "R55,990", "₦4,630,000", "KSh 449,990", "GH₵ 37,990", "$2,990"],
    ]
    return table(s, MX, y, [3.63, 1.7, 1.7, 1.7, 1.7, 1.7], [0.62, 0.66, 0.66, 0.66], rows, size=16, label="wholesale")


def p_wholesale(prs, v, label=None):
    s = new_slide(prs)
    eyebrow = f"Appendix {label} · partner wholesale" if label else "30% off Agency list · every billing currency"
    header(s, eyebrow, "Partner wholesale pricing")
    wholesale_table(s, 1.65)
    g = gid("illus")
    box(s, MX, 4.62, 7.35, 2.05, fill=MIST, shape="round", radius=0.18, role="card", group=g)
    text(s, MX + 0.35, 4.82, 6.6, 0.3, "ILLUSTRATIVE ONLY · NOT A FORECAST", size=12, bold=True, color=DEEP,
         role="eyebrow", group=g, spc=120)
    text(s, MX + 0.35, 5.2, 6.7, 1.3, [{"runs": "20 clients × R1,500/mo = R30,000 billed", "space_after": 6},
                                       {"runs": "− R5,599 wholesale = R24,401/mo gross spread, before your own costs",
                                        "size": 16, "bold": False}],
         size=19, font=HEAD, bold=True, role="card", group=g)
    text(s, 8.3, 4.72, 4.43, 1.9, [{"runs": "Same inclusions as Agency. Annual = 10× monthly.", "space_after": 8},
                                   {"runs": "You set your retail price.", "space_after": 8},
                                   "No payment rail? Waitlist only."],
         size=14, color=MUTED, role="body")
    notes(s, """Partner wholesale pricing is the Agency tier at 30% off list, in every currency we bill: R5,599,
    ₦463,000, KSh 44,999, GH₵ 3,799 or $299 a month, and ten times that for a year. It includes everything in
    Agency: full white-label, 50 client sub-accounts, reseller billing, 80 channels, 500,000 AI credits and dedicated
    support. Launch offers don't stack with wholesale. The example is illustrative only: a partner sets
    their own retail price and carries their own delivery costs. Referral commission on clients' own plans is still
    a founder decision.""", v)
    return s


def p_onboarding(prs, v):
    s = new_slide(prs)
    header(s, "Becoming a partner", "From first call to first client")
    steps = [("search", "Intro call", "Clients, services, pricing"),
             ("sparkles", "White-label demo", "FluxMuse under your brand"),
             ("target", "Pilot client", "Launch one client first"),
             ("briefcase", "Go wholesale", "Agency tier at R5,599/mo"),
             ("megaphone", "Roll out", "Onboard more clients")]
    colw = 2.33
    gapx = (CW - 5 * colw) / 4
    box(s, MX + colw / 2, 2.2, CW - colw, 0.05, fill=ORANGE, role="deco", group="line")
    for i, (ic, t, sub) in enumerate(steps):
        x = MX + i * (colw + gapx)
        cx = x + colw / 2
        g = gid("step")
        box(s, cx - 0.4, 1.83, 0.8, 0.8, fill=PAPER, line=ORANGE, lw=2.25, shape="oval", role="deco", group=g)
        text(s, cx - 0.4, 1.83, 0.8, 0.8, str(i + 1), size=20, font=HEAD, bold=True, color=DEEP, align="c",
             anchor="m", role="card", group=g)
        box(s, x, 2.95, colw, 3.2, fill=MIST, shape="round", radius=0.18, role="card", group=g)
        icon(s, ic, "deep", x + 0.6, 3.6, 0.72, circle=PAPER, group=g)
        text(s, x + 0.25, 4.2, colw - 0.45, 0.8, t, size=18, font=HEAD, bold=True, role="card", group=g)
        text(s, x + 0.25, 4.95, colw - 0.45, 1.1, sub, size=14, color=MUTED, role="card", group=g)
    text(s, MX, 6.4, CW, 0.3, "Timings depend on your team. Partner terms and co-marketing: [[TBC]].", size=12,
         color=MUTED, role="caption")
    notes(s, """Becoming a partner follows the same path as our agency buying motion. It starts with an intro call
    about your clients, services and pricing. Then a white-label demo so you can see the platform under your brand.
    Next, pick one pilot client and launch them first, so you learn the workflow on a real account. Once that's
    working, move onto the Agency tier at wholesale pricing with reseller billing, and roll out to more clients. We
    don't quote fixed onboarding timings yet, and the formal partner terms are still being confirmed.""", v)
    return s


def p_comarketing(prs, v):
    s = new_slide(prs)
    header(s, "Co-marketing & commitments", "What we bring, and what we ask")
    cols = [("FluxMuse provides", TINT, "check",
             ["Partner directory listing", "This deck and proposal templates",
              "Consented pilot case studies (Dec 2026)", "Co-marketing support: [[TBC]]"]),
            ("Partners commit to", MIST, "file",
             ["Partner obligations: [[TBC]]", "Minimum active clients: [[TBC]]",
              "Brand and POPIA standards: [[TBC]]",
              "Referral commission: [[FOUNDER DECISION: referral commission %]]"])]
    for c, (head, fill, ic, items) in enumerate(cols):
        x = MX + c * (5.95 + 0.233)
        g = gid("col")
        box(s, x, 1.68, 5.95, 5.05, fill=fill, shape="round", radius=0.2, role="card", group=g)
        text(s, x + 0.4, 1.98, 5.1, 0.5, head, size=22, font=HEAD, bold=True, role="card", group=g)
        for i, t in enumerate(items):
            y = 2.75 + i * 0.95
            gg = icon(s, ic, "deep", x + 0.68, y + 0.36, 0.56, circle=PAPER)
            role = "placeholder" if "[[" in t else "body"
            text(s, x + 1.15, y, 4.5, 0.74, t, size=16, anchor="m", role=role, group=gg)
    notes(s, """Be transparent about what's defined and what isn't. FluxMuse provides a partner directory in the
    product, sales material such as this deck and proposal templates, and consented pilot case studies from
    December 2026. Co-marketing specifics are still being finalised. Partner obligations, such as minimum active
    clients, brand standards and data-protection responsibilities under POPIA, are to be confirmed. Referral
    commission for clients who sign up to their own Starter, Growth or Scale plans is a founder decision that hasn't
    been made yet. Don't promise terms; take their questions and follow up in writing.""", v)
    return s


# ---------------------------------------------------------------- appendix
def a_divider(prs, v, items):
    s = new_slide(prs, dark=True)
    text(s, MX, 1.5, CW, 0.3, "APPENDIX", size=12, bold=True, color=ORANGE, role="eyebrow", spc=150)
    text(s, MX, 1.82, CW, 0.9, "Reference slides", size=40, font=HEAD, bold=True, color=PAPER, role="title")
    for i, (lab, t) in enumerate(items):
        y = 3.05 + i * 0.78
        g = gid("ax")
        box(s, MX, y, 0.95, 0.56, fill=ORANGE, shape="round", radius=0.28, role="chip", group=g)
        text(s, MX, y, 0.95, 0.56, lab, size=16, font=HEAD, bold=True, color=NIGHT, align="c", anchor="m",
             role="chip", group=g)
        text(s, MX + 1.2, y, 8, 0.56, t, size=20, color=PAPER, anchor="m", role="body", group=g)
    notes(s, """The appendix is for questions and follow-up, not the main talk track. It covers local-currency pricing
    for Nigeria, Kenya and Ghana plus USD for the other rail-covered countries, the full matrix of payment rails by
    country, partner wholesale pricing where relevant, and the questions prospects ask most. Only show prices to
    prospects in countries we sell to. If someone is in a country without a payment rail, including Botswana and
    Namibia for now, share the waitlist link instead of prices: [[WAITLIST LINK]].""", v)
    return s


def a_regional_img(prs, v, label):
    s = new_slide(prs)
    header(s, f"Appendix {label} · Nigeria, Kenya, Ghana & USD markets", "Regional pricing")
    img(s, "regional", "Monthly plan prices in ZAR, NGN, KES, GHS and USD")
    notes(s, """Use this slide only for prospects in Nigeria, Kenya, Ghana or a USD market. It's one set of plans at
    fixed local price points, set at FX parity with the rand in early September 2026 and reviewed quarterly. Nigeria,
    Kenya and Ghana are billed in naira, shillings and cedis; the other 19 rail-covered countries are billed in US
    dollars. Annual plans are ten times monthly. The live website may still show FX-converted amounts, so quote this
    slide and say "billed in local currency". Everyone else joins the waitlist.""", v)
    return s


def a_regional_table(prs, v, label):
    s = new_slide(prs)
    header(s, f"Appendix {label} · editable", "Local-currency prices, monthly and annual")
    hdr = {"fill": NIGHT, "color": PAPER, "bold": True, "font": HEAD, "size": 14, "line": None}
    rows = [[dict(hdr, text="Plan", align="l")] + [dict(hdr, text=t) for t in
             ["NGN / mo", "NGN / yr", "KES / mo", "KES / yr", "GHS / mo", "GHS / yr", "USD / mo", "USD / yr"]]]
    data = [("Starter", "₦41,000", "₦410,000", "KSh 3,999", "KSh 39,990", "GH₵ 339", "GH₵ 3,390", "$27", "$270"),
            ("Growth", "₦165,000", "₦1,650,000", "KSh 15,999", "KSh 159,990", "GH₵ 1,359", "GH₵ 13,590", "$109", "$1,090"),
            ("Scale", "₦413,000", "₦4,130,000", "KSh 39,999", "KSh 399,990", "GH₵ 3,399", "GH₵ 33,990", "$269", "$2,690"),
            ("Agency", "₦662,000", "₦6,620,000", "KSh 64,499", "KSh 644,990", "GH₵ 5,439", "GH₵ 54,390", "$429", "$4,290"),
            ("Enterprise", "from ₦1,650,000", "custom", "from KSh 161,000", "custom", "from GH₵ 13,600", "custom",
             "from $1,099", "custom")]
    for r in data:
        fill = TINT if r[0] == "Growth" else PAPER
        rows.append([{"text": r[0], "bold": True, "fill": fill}] + [{"text": t, "fill": fill} for t in r[1:]])
    table(s, MX, 1.65, [1.33] + [1.35] * 8, [0.62] + [0.68] * 5, rows, size=14, label="regional")
    text(s, MX, 5.95, CW, 0.75, [{"runs": "Billed in local currency at fixed price points, reviewed quarterly. USD "
                                          "applies in 19 other rail-covered countries.", "space_after": 4},
                                 "Countries without a payment rail, including Botswana and Namibia (coming soon): "
                                 "waitlist only, no prices."],
         size=12, color=MUTED, role="caption")
    notes(s, """Same prices as the previous slide, as an editable table with annual prices added. Starter is ₦41,000,
    KSh 3,999, GH₵ 339 or $27 a month. Growth is ₦165,000, KSh 15,999, GH₵ 1,359 or $109. Scale is ₦413,000, KSh
    39,999, GH₵ 3,399 or $269. Agency is ₦662,000, KSh 64,499, GH₵ 5,439 or $429. Enterprise starts from ₦1,650,000,
    KSh 161,000, GH₵ 13,600 or $1,099. If a prospect asks about another currency, the answer is USD, and never show
    these prices to a country we don't sell to.""", v)
    return s


def a_rails(prs, v, label):
    s = new_slide(prs)
    header(s, f"Appendix {label} · payment rails", "Payment rails by country")
    img(s, "matrix", "Matrix of 23 countries by region against 5 payment rails with billing currency; Botswana and "
        "Namibia coming soon")
    notes(s, """This matrix shows which of the five rails covers each of the 23 countries, grouped by region, with the
    billing currency. South Africa has four rails: Yoco, Ozow, Paystack and Fincra. Nigeria, Ghana and Kenya each
    have Paystack, pawaPay and Fincra. Most other countries are covered by pawaPay mobile money and Fincra. South
    Sudan and Zimbabwe are Fincra payouts only, so sales there open once subscription collection is confirmed.
    Botswana and Namibia are coming soon with no rail yet, so they're waitlist only.""", v)
    return s


FAQ = [
    ("Is my customers' data safe?",
     "FluxMuse is POPIA, NDPR and GDPR ready, with row-level security on every table, encrypted tokens, audit export, "
     "and account deletion and data export."),
    ("Can I use my own WhatsApp number?",
     "You connect a WhatsApp Business number through Meta's official Embedded Signup (WhatsApp Cloud API). We check "
     "your number with you during setup."),
    ("Which payment methods can customers use?",
     "In South Africa: Yoco card and tap-to-pay, Ozow instant EFT and Paystack. Across Africa: Paystack, pawaPay "
     "mobile money and Fincra, in 23 countries."),
    ("Is there a contract? Can I cancel?",
     "Start with a 14-day free trial, no card required. Then pay monthly, or annually for 2 months free. "
     "Cancellation terms: [[CONFIRM TERMS]]."),
    ("Which languages does it write in?",
     "The Creator agent writes in isiZulu, Afrikaans, English, Pidgin, Swahili and more."),
    ("How long does onboarding take?",
     "Typically 1–2 weeks to set up your WhatsApp number, catalog, payment rail and content calendar."),
    ("I already use Shopify or WooCommerce.",
     "Keep them. FluxMuse integrates with Shopify, WooCommerce and Takealot, plus HubSpot/CRM sync, accounting sync, "
     "Slack and Zapier, from the Growth plan."),
    ("Where can I buy FluxMuse?",
     "In 23 African countries with local payment rails: ZAR, NGN, KES or GHS pricing, USD elsewhere. Botswana and "
     "Namibia are coming soon; other countries can join the waitlist."),
]


def a_faq(prs, v, label, part):
    s = new_slide(prs)
    items = FAQ[:4] if part == 1 else FAQ[4:]
    header(s, f"Appendix {label} · FAQ ({part} of 2)", "Questions prospects ask")
    cwid, ch = 5.95, 2.45
    for i, (q, a) in enumerate(items):
        x = MX + (i % 2) * (cwid + 0.233)
        y = 1.68 + (i // 2) * (ch + 0.2)
        g = gid("faq")
        box(s, x, y, cwid, ch, fill=MIST, shape="round", radius=0.18, role="card", group=g)
        icon(s, "chat", "deep", x + 0.55, y + 0.58, 0.6, circle=PAPER, group=g)
        text(s, x + 1.0, y + 0.2, cwid - 1.3, 0.78, q, size=17, font=HEAD, bold=True, anchor="m", role="card", group=g)
        role = "placeholder" if "[[" in a else "card"
        text(s, x + 0.35, y + 1.08, cwid - 0.7, 1.24, a, size=15, role=role, group=g)
    if part == 1:
        n = """Use these answers as written; they're checked against what the product does today. On data, stick to
        the controls listed and don't claim certifications we haven't confirmed. On WhatsApp numbers, explain that the
        number is connected through Meta's Embedded Signup and that we check it together during setup, because a
        number already used in the WhatsApp Business app may need extra steps. On payments, name only the five rails.
        On contracts, cancellation terms are still to be confirmed, so don't improvise; promise a written answer."""
    else:
        n = """Languages: the Creator agent writes in isiZulu, Afrikaans, English, Pidgin, Swahili and more, so offer to
        generate a sample post in their language during the trial. Onboarding: typical setup takes one to two weeks,
        but Meta template approvals and catalog size can stretch that, so set expectations honestly. Shopify and
        WooCommerce users keep their store, and FluxMuse connects to it from the Growth plan. Availability: we sell in
        23 countries with payment rails; anyone in Botswana, Namibia or another gated country gets the waitlist link,
        never prices."""
    notes(s, n, v)
    return s


# ================================================================ DECKS
def deck_defs():
    B = lambda f, **kw: (f, kw)
    return {
        "FluxMuse_Brand_Pitch_Deck.pptx": ("master", [
            B(s_title), B(s_reality), B(s_meet), B(s_how), B(s_team), B(s_sell), B(s_channels), B(s_paid),
            B(s_segments), B(s_seg_solo), B(s_seg_sme), B(s_seg_agency), B(s_partner_model), B(s_action),
            B(s_journey), B(s_pilot), B(s_case), B(s_pricing), B(s_pricing_table), B(s_founding), B(s_trust),
            B(s_roadmap), B(s_next),
            B(a_divider, items=[("A1", "Regional pricing: Nigeria, Kenya, Ghana & USD markets"),
                                ("A2", "Payment rails by country"), ("A3", "Partner wholesale pricing"),
                                ("A4", "Frequently asked questions")]),
            B(a_regional_img, label="A1"), B(a_regional_table, label="A1"), B(a_rails, label="A2"),
            B(p_wholesale, label="A3"), B(a_faq, label="A4", part=1), B(a_faq, label="A4", part=2),
        ]),
        "FluxMuse_Pitch_Solo_Entrepreneurs.pptx": ("solo", [
            B(s_title), B(s_reality), B(s_meet), B(s_how), B(s_team), B(s_sell), B(s_paid), B(s_seg_solo),
            B(s_action), B(s_case), B(s_pricing), B(s_founding), B(s_trust), B(s_next),
            B(a_faq, label="A1", part=1), B(a_faq, label="A1", part=2),
        ]),
        "FluxMuse_Pitch_SMEs.pptx": ("sme", [
            B(s_title), B(s_reality), B(s_meet), B(s_how), B(s_team), B(s_sell), B(s_channels),
            B(s_sme_integrations), B(s_paid), B(s_seg_sme), B(s_action), B(s_journey), B(s_pilot), B(s_pricing),
            B(s_pricing_table), B(s_founding), B(s_trust), B(s_roadmap), B(s_next),
            B(a_divider, items=[("A1", "Regional pricing: Nigeria, Kenya, Ghana & USD markets"),
                                ("A2", "Payment rails by country"), ("A3", "Frequently asked questions")]),
            B(a_regional_img, label="A1"), B(a_regional_table, label="A1"), B(a_rails, label="A2"),
            B(a_faq, label="A3", part=1), B(a_faq, label="A3", part=2),
        ]),
        "FluxMuse_Partner_Programme_Deck.pptx": ("partner", [
            B(s_title), B(p_squeeze), B(s_meet), B(s_how), B(s_team), B(s_sell), B(s_channels), B(s_seg_agency),
            B(s_partner_model), B(p_what_you_get), B(p_wholesale), B(p_onboarding), B(p_comarketing), B(s_paid),
            B(s_pilot), B(s_trust), B(s_roadmap), B(s_next),
            B(a_divider, items=[("A1", "Regional pricing: Nigeria, Kenya, Ghana & USD markets"),
                                ("A2", "Payment rails by country"), ("A3", "Frequently asked questions")]),
            B(a_regional_img, label="A1"), B(a_regional_table, label="A1"), B(a_rails, label="A2"),
            B(a_faq, label="A3", part=1), B(a_faq, label="A3", part=2),
        ]),
    }


def set_theme_fonts(prs):
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            tp = rel.target_part
            blob = tp.blob
            new = re.sub(rb'(<a:majorFont>\s*<a:latin typeface=")[^"]*"', rb'\1Poppins"', blob)
            new = re.sub(rb'(<a:minorFont>\s*<a:latin typeface=")[^"]*"', rb'\1Inter"', new)
            if hasattr(tp, "_element"):
                tp._element = etree.fromstring(new)
            else:
                tp._blob = new


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    marker = V3_MARKER.exists()
    manifest = {"built": date.today().isoformat(), "assets_revision_v3_marker": marker, "decks": {}}
    for fname, (v, slides) in deck_defs().items():
        _CUR_DECK[0] = fname
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
        set_theme_fonts(prs)
        prs.core_properties.title = "FluxMuse pitch deck" if v != "partner" else "FluxMuse Partner Programme"
        prs.core_properties.author = "Fluxmuse Pty Ltd"
        prs.core_properties.subject = "AI marketing & WhatsApp commerce"
        for i, (fn, kw) in enumerate(slides):
            s = fn(prs, v, **kw)
            if i > 0:
                chrome(s, i + 1, getattr(s, "_fm_dark", False))
        prs.save(OUT / fname)
        manifest["decks"][fname] = {"variant": v, "slides": len(prs.slides),
                                    "images": sorted(USED_IMAGES.get(fname, []))}
        print(f"built {fname}: {len(prs.slides)} slides")
    (HERE / "build_manifest.json").write_text(json.dumps(manifest, indent=2))
    print(f"assets/.revision-v3-done present: {marker}")


if __name__ == "__main__":
    build()
