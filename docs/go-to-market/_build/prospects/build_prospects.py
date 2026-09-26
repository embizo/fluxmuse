#!/usr/bin/env python3
"""Builds the prospect pack from docs/go-to-market/08_Prospects/data/*.json.

  python build_prospects.py            validate + build everything
  python build_prospects.py --check    validate only
  python build_prospects.py --selftest render a synthetic record to a temp dir (no data needed)

Outputs (docs/go-to-market/08_Prospects/):
  proposals/NN_<slug>.docx   client-facing proposal, one per prospect
  Prospect_Briefs.docx       INTERNAL: approach route, opening message, fit and gaps, one page each
  Prospect_List.xlsx         INTERNAL: master list, pipeline tracker, sources, pipeline maths

Offer, prices and policy come from 08_Prospects/CURRENT_OFFER.md. If a price or policy changes, edit the
constants below, not the JSON. Reuses the house-style document class from ../proposals/fmdoc.py.
"""
import argparse
import json
import re
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "proposals"))
from docx.enum.text import WD_TAB_ALIGNMENT  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
from fmdoc import (ASSETS, DEEP, GREY, HEAD_FONT, PPR_ORDER, SLATE, TINT, FMDoc, border_el,  # noqa: E402
                   img_stream, put)

GTM = HERE.parent.parent
OUT = GTM / "08_Prospects"
DATA = OUT / "data"
AS_AT = "26 September 2026"

# ---- constants that mirror CURRENT_OFFER.md (edit here when the offer changes) ------------------------
SHOW_FOUNDING_MEMBER = True          # one optional line; re-confirm the promotion is still enabled
GUARANTEE_TERMS = "[[GUARANTEE TERMS: money-back guarantee, terms to be confirmed by the founder]]"
CHECKOUT_FEE_WORDING = "[[FOUNDER TO CONFIRM CHECKOUT FEE WORDING]]"
VAT_NOTE = "Prices are in South African rand (ZAR). VAT treatment: [[CONFIRM VAT TREATMENT]]."
TIERS = {  # code: (display name, monthly ZAR, annual ZAR, what the plan includes)
    "nano": ("Nano", 149, 1490, "1 brand, 3 channels"),
    "micro": ("Micro", 289, 2890, "1 brand, 4 channels, 1 WhatsApp number"),
    "starter": ("Starter", 499, 4990, "1 brand, 3 channels"),
    "growth": ("Growth", 1999, 19990, "3 brands, 15 channels, inbound AI Voice (beta)"),
    "scale": ("Scale", 4999, 49990, "10 brands, 40 channels, inbound and outbound AI Voice (beta)"),
    "corporate": ("Corporate", 6999, 69990, "25 brands, 60 channels, bring-your-own-cloud"),
    "agency": ("Agency", 9999, 99990, "unlimited brands, 80 channels, white-label, bring-your-own-cloud, multi-client"),
    "custom": ("Custom", None, None, "scoped together, priced by consultation"),
}
FOUNDING_TIERS = {"starter", "growth", "scale"}
STATUS_LABEL = {"live": "Live (newly launched)", "in_setup": "Being switched on"}
SEGMENTS = {"creator", "solo", "sme", "agency"}

# ---- validation ------------------------------------------------------------------------------------------
FORBIDDEN = re.compile(
    r"free trial|free month|free period|first month (is )?free|free first month|14[- ]day (free )?trial|14 days free|"
    r"try (it )?free|risk[- ]free|\bpilot\b|our customers|trusted by|\bproven\b|"
    r"guarantee[sd]? (results|sales|roi|leads|revenue)|\b10x\b|\bdiscount|sales uplift|snapscan|stitch\b|"
    r"flutterwave|payfast|m-pesa direct|money[- ]back|\\brefund|no lock[- ]in|no contract|cancel anytime|"
    r"(tell us|let us know) and we (stop|refund|cancel)|we stop\\b|nothing to lose|are yours|belong to you|"
    r"we drop it|reason not to continue",
    re.I)
# Fields where WE make offers or promises. Facts about the prospect (public_facts, evidence, gaps) are not scanned.
OFFER_KEYS = ("one_line_pitch", "why_wave", "recommended_plan", "what_we_would_do", "first_30_days", "we_would_track",
              "objections", "fit")
OPT_OUT = re.compile(
    r"opt[- ]?out|\bstop\b|don'?t want|not interested|not relevant|rather (we|i) (didn'?t|not|don'?t)|"
    r"let (me|us) know if you'?d rather|no need to reply|unsubscribe|just say so|tell (me|us) and (i|we)|"
    r"(won'?t|will not|shall not) (follow up|contact|message|bother|chase|write)|rather not|not hear from us|"
    r"reply ['\"]?no", re.I)
THIRD = re.compile(r"\b(his|her|hers|him|he|she|the owner|the owners|the founder|his team|her team|the business's)\b", re.I)
CLIENT_TEXT_KEYS = ("one_line_pitch", "what_we_would_do", "first_30_days", "we_would_track", "objections",
                    "recommended_plan")
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?<!\d)(?:\+27|0)[\s-]?\d{2}[\s-]?\d{3}[\s-]?\d{4}(?!\d)")
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _strings(o):
    if isinstance(o, str):
        yield o
    elif isinstance(o, dict):
        for v in o.values():
            yield from _strings(v)
    elif isinstance(o, list):
        for v in o:
            yield from _strings(v)


def validate(rec, fname=""):
    """Return (errors, warnings). Errors block the build; warnings need a human look."""
    E, W = [], []
    need = ["id", "slot", "name", "segment", "trade", "location", "wave", "why_wave", "one_line_pitch", "public_facts",
            "decision_maker", "approach", "pains", "fit", "recommended_plan", "what_we_would_do", "first_30_days",
            "we_would_track", "objections", "gaps"]
    for k in need:
        if k not in rec:
            E.append(f"missing key: {k}")
    if E:
        return E, W
    if rec["segment"] not in SEGMENTS:
        E.append(f"segment must be one of {sorted(SEGMENTS)}")
    if rec["wave"] not in (1, 2, 3):
        E.append("wave must be 1, 2 or 3")
    if not 6 <= len(rec["public_facts"]) <= 10:
        E.append(f"public_facts must be 6-10 (has {len(rec['public_facts'])})")
    for i, f in enumerate(rec["public_facts"]):
        if not str(f.get("source", "")).startswith("http"):
            E.append(f"public_facts[{i}] has no http source")
        if not DATE.match(str(f.get("retrieved", ""))):
            E.append(f"public_facts[{i}] retrieved must be YYYY-MM-DD")
    if not 3 <= len(rec["pains"]) <= 5:
        E.append(f"pains must be 3-5 (has {len(rec['pains'])})")
    for i, p in enumerate(rec["pains"]):
        if p.get("basis") not in ("observed", "inferred"):
            E.append(f"pains[{i}].basis must be observed|inferred")
        for k in ("pain", "evidence", "validate_with"):
            if not p.get(k):
                E.append(f"pains[{i}].{k} empty")
    if not 4 <= len(rec["what_we_would_do"]) <= 6:
        E.append(f"what_we_would_do must be 4-6 (has {len(rec['what_we_would_do'])})")
    for i, c in enumerate(rec["what_we_would_do"]):
        if c.get("status") not in STATUS_LABEL:
            E.append(f"what_we_would_do[{i}].status must be live|in_setup")
        if re.search(r"instagram|facebook", c.get("capability", "") + c.get("for_them", ""), re.I) and c.get("status") == "live" \
                and re.search(r"publish|post|dm|comment", c.get("capability", "") + c.get("for_them", ""), re.I):
            W.append(f"what_we_would_do[{i}] marks Instagram/Facebook posting as live: check CURRENT_OFFER §3")
    if [w.get("week") for w in rec["first_30_days"]] != [1, 2, 3, 4]:
        E.append("first_30_days must be weeks 1,2,3,4")
    if not 3 <= len(rec["we_would_track"]) <= 5:
        E.append("we_would_track must be 3-5")
    if len(rec["objections"]) != 3:
        E.append("objections must be exactly 3")
    if rec["recommended_plan"].get("tier") not in TIERS:
        E.append(f"recommended_plan.tier must be one of {sorted(TIERS)}")
    a = rec["approach"]
    for k in ("route", "contact_source", "opening_message", "compliance_note"):
        if not a.get(k):
            E.append(f"approach.{k} empty")
    words = len(a.get("opening_message", "").split())
    if not 40 <= words <= 120:
        W.append(f"opening_message is {words} words (aim 60-90)")
    if not OPT_OUT.search(a.get("opening_message", "")):
        E.append("opening_message has no opt-out line")
    if not 1 <= int(rec["fit"].get("attainability", 0)) <= 5:
        E.append("fit.attainability must be 1-5")
    for s in _strings({k: rec[k] for k in OFFER_KEYS if k in rec}):
        m = FORBIDDEN.search(s)
        if m:
            E.append(f"forbidden phrase '{m.group(0)}' in: {s[:90]}")
    m = FORBIDDEN.search(a.get("opening_message", ""))
    if m:
        E.append(f"forbidden phrase '{m.group(0)}' in opening_message")
    for s in _strings({k: rec[k] for k in CLIENT_TEXT_KEYS if k in rec}) :
        m = THIRD.search(s)
        if m:
            W.append(f"third person '{m.group(0)}' in client-facing text (should say 'you'): {s[:80]}")
    for s in _strings({"facts": rec["public_facts"], "channels": rec.get("channels_observed", []), "approach": a}):
        if EMAIL.search(s) or PHONE.search(s):
            W.append(f"contact detail present, confirm it is published on the business's own page: {s[:80]}")
    return E, W


# ---- documents -------------------------------------------------------------------------------------------
class ProspectDoc(FMDoc):
    """FMDoc with a per-prospect header and no proposal-number placeholder."""

    footer_text = "Fluxmuse Pty Ltd · fluxmuse.ai · Confidential"

    def __init__(self, header_right, client_ph):
        self.header_right = header_right
        super().__init__(sample=False, client_ph=client_ph)

    def _header_para(self, header):
        p = header.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
        put(p._p.get_or_add_pPr(), border_el("w:pBdr", ["bottom"], sz=12, color="FF6A00", space=4), PPR_ORDER)
        r = p.add_run()
        r.add_picture(img_stream(ASSETS / "brand/fluxmuse-icon-512.png", 256), height=Cm(0.8))
        self._docpr(r, "FluxMuse icon")
        self._run(p, "  FLUX", font=HEAD_FONT, size=12, color=DEEP, bold=True)
        self._run(p, "MUSE", font=HEAD_FONT, size=12, color=SLATE, bold=True)
        self._run(p, "\t" + self.header_right, size=8, color=GREY)

    def _footer_para(self, footer):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        put(p._p.get_or_add_pPr(), border_el("w:pBdr", ["top"], sz=4, color="D9DDE2", space=6), PPR_ORDER)
        self.rich(p, self.footer_text + " · Page ", size=8, color=GREY)
        self._field(p, "PAGE")
        self.rich(p, " of ", size=8, color=GREY)
        self._field(p, "NUMPAGES")

    def cover_page(self, eyebrow, title, subtitle, rows, note):
        self.spacer(10)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run()
        r.add_picture(img_stream(ASSETS / "brand/fluxmuse-icon-512.png", 512), height=Cm(3.2))
        self._docpr(r, "FluxMuse Muse icon: a multicolour low-poly head over a node network")
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        self._run(p, "FLUX", font=HEAD_FONT, size=30, color=DEEP, bold=True)
        self._run(p, "MUSE", font=HEAD_FONT, size=30, color=SLATE, bold=True)
        self.para("AI Marketing & WhatsApp Commerce for Africa", size=11, color=GREY, after=26)
        self.eyebrow(eyebrow)
        p = self.doc.add_paragraph(style="Title")
        p.paragraph_format.space_after = Pt(4)
        self.rich(p, title)
        self.para(subtitle, size=13, color=SLATE, after=14)
        self.rule("FF6A00", 24, after=10)
        self.kv_table(rows, caption="Layout: proposal details")
        self.para(note, size=8.5, color=GREY, before=10)
        self.page_break()


ROUTE_FACT = re.compile(r"\b(e-?mail address(es)?|phone numbers?|contact page|contact form|contact details|"
                        r"bookings? e-?mail|press e-?mail|PR/interview)\b", re.I)


def client_facts(rec):
    """Facts safe to show back to the business: drops any that are about HOW TO CONTACT them (route facts stay in the
    internal brief) or that quote their own phone number or email."""
    return [f for f in rec["public_facts"]
            if not (EMAIL.search(f["fact"]) or PHONE.search(f["fact"]) or ROUTE_FACT.search(f["fact"]))]


def source_index(rec):
    """Facts shown to the client, plus the de-duplicated, numbered list of URLs they cite."""
    shown = client_facts(rec)
    urls = []
    for f in shown:
        if f["source"] not in urls:
            urls.append(f["source"])
    for b in rec.get("businesses", []):
        if b.get("source") and b["source"] not in urls:
            urls.append(b["source"])
    return shown, urls


def money(n):
    return f"R{n:,.0f}"


def plan_rows(rec):
    code = rec["recommended_plan"]["tier"]
    name, monthly, annual, incl = TIERS[code]
    rows = [("Recommended plan", f"**{name}**" + (f": {money(monthly)} per month" if monthly else ": priced after our call"))]
    rows.append(("Includes", incl))
    if annual:
        rows.append(("Annual option", f"{money(annual)} for 12 months (10 times the monthly price)"))
    rows.append(("Why this plan", rec["recommended_plan"]["why"]))
    rows.append(("When to move up", rec["recommended_plan"]["upgrade_path"]))
    rows.append(("How payment works", "Payment is taken first, when you subscribe. " + GUARANTEE_TERMS))
    if SHOW_FOUNDING_MEMBER and code in FOUNDING_TIERS:
        rows.append(("South African launch pricing",
                     "30% off your first two monthly billing cycles is applied at checkout while the offer is open. "
                     "[[CONFIRM the offer is still active before sending]]"))
    extra = "Meta charges for some WhatsApp conversations, and AI images and video use your monthly AI credits. " \
            "We will show you both before you pay."
    if any("checkout" in (c["capability"] + c["for_them"]).lower() for c in rec["what_we_would_do"]):
        extra += " Checkout through FluxMuse carries a per-order fee: " + CHECKOUT_FEE_WORDING
    rows.append(("Other costs", extra))
    return rows


def build_proposal(rec, outdir):
    n = rec["name"]
    d = ProspectDoc(f"Proposal for {n}", n)
    d.footer_text = f"Fluxmuse Pty Ltd · fluxmuse.ai · Confidential proposal for {n}"
    d.cover_page("Proposal", f"A WhatsApp-first plan for {n}", f"{rec['trade']} · {rec['location']}",
                 [("Prepared for", n), ("Date", f"{AS_AT} (draft)"),
                  ("Prepared by", "[[SENDER NAME]], [[SENDER TITLE]], Fluxmuse Pty Ltd\n[[SENDER EMAIL]] · [[SENDER WHATSAPP]]"),
                  ("Valid for", "30 days from the date it is sent")],
                 f"Confidential. Prepared for {n} only. {VAT_NOTE}")

    d.h1("1. Why we are writing to you")
    d.callout("In one sentence", [rec["one_line_pitch"]], fill=TINT, accent="FF6A00")
    d.para("FluxMuse is a new AI marketing and WhatsApp commerce platform built for African small businesses. "
           "We will be plain about where we are: we have no customer numbers or results to show you yet. We are "
           "opening with a small first group of businesses and setting each one up by hand. What we can offer is a "
           "working product, set up with you, and the tracking to show what it does for your business.")

    d.h1("2. What we noticed about " + n)
    d.para(f"This comes from your own public pages and press, as at {AS_AT}. Sources are listed at the end. "
           "If anything is out of date or wrong, tell us and we will correct it.", size=9.5, color=GREY)
    shown, urls = source_index(rec)
    facts = [(f["fact"], f"[{urls.index(f['source']) + 1}]") for f in shown]
    d.table(["What we found", "Source"], facts, [14.4, 2.6], "What we found about the business", size=9.5,
            first_col_bold=False)
    if rec.get("businesses"):
        d.h3("The businesses we looked at")
        d.table(["Business", "What it is"], [(b["name"], b["what"]) for b in rec["businesses"]], [6.0, 11.0],
                "Businesses", size=9.5)

    d.h1("3. Where time and orders may be leaking")
    d.para("These are our hypotheses, not findings. The first call is about testing them with you. "
           "If we have any of them wrong, that is useful for both of us.", size=9.5, color=GREY)
    rows = []
    for p in rec["pains"]:
        how = ("Seen on your public pages: " if p["basis"] == "observed" else "Typical for this kind of business: ") + p["evidence"]
        rows.append((p["pain"], how, p["validate_with"]))
    d.table(["Where it may be leaking", "Why we think so", "What we would ask you"], rows, [5.6, 5.9, 5.5],
            "Hypotheses about where time and orders leak", size=9)

    d.h1("4. What FluxMuse would do for you")
    caps = [(c["capability"], STATUS_LABEL[c["status"]], c["for_them"]) for c in rec["what_we_would_do"]]
    d.table(["What", "Status", "What it means for you"], caps, [4.6, 3.0, 9.4], "What FluxMuse would do", size=9.5)
    d.para("**Status key.** Live (newly launched) means built and running, but no outside business has run it end to end "
           "yet, so we would watch it closely with you. Being switched on means it depends on an approval or final "
           "set-up, and we will not count on it until it is confirmed.", size=9, color=GREY)

    d.h1("5. Your first 30 days")
    d.table(["Week", "What happens"],
            [(f"Week {w['week']}", "\n".join("• " + a for a in w["actions"])) for w in rec["first_30_days"]],
            [2.4, 14.6], "First 30 days", size=9.5)

    d.h1("6. What we would track together")
    d.bullets(rec["we_would_track"])
    d.para("We agree the numbers with you on the first call. We do not promise a number of leads or sales.",
           size=9.5, color=GREY)

    d.h1("7. Your plan and price")
    d.kv_table(plan_rows(rec), widths=(4.4, 12.6), caption="Layout: plan and price", size=9.5)
    d.para(VAT_NOTE, size=8.5, color=GREY)

    d.h1("8. Questions you may have")
    d.table(["A question you might have", "Our honest answer"],
            [(o["objection"], o["response"]) for o in rec["objections"]], [6.0, 11.0], "Questions and answers", size=9.5)
    if rec.get("related_reading"):
        d.para("**Further reading**", after=2, keep=True)
        d.bullets([f"{r['title']}: {r['url']}" for r in rec["related_reading"]])

    d.h1("9. Next step")
    d.para("A 20-minute call to walk through this and decide whether it is worth doing. No obligation.")
    d.kv_table([("Who", "[[SENDER NAME]], Fluxmuse Pty Ltd"), ("Email", "[[SENDER EMAIL]]"),
                ("WhatsApp", "[[SENDER WHATSAPP]]")], caption="Layout: contact", size=10)
    d.para("If you would rather we did not contact you again, tell us and we will stop.", size=9.5, color=GREY)

    d.h1("10. About FluxMuse")
    d.bullets([
        "**Fluxmuse Pty Ltd**, South Africa. A verified Meta Tech Provider (business and access verification, as at "
        "September 2026).",
        "**Built for African small businesses:** WhatsApp first, priced in rand, and able to answer in your customer's language.",
        "**Privacy:** messages go only to people who opted in, every broadcast has an opt-out, and you can export or "
        "delete your data. [[LEGAL REVIEW REQUIRED before sending: POPIA wording]]",
        "**Newly launched:** several features have not yet been run end to end by an outside business. We say so above, "
        "feature by feature.",
    ])

    d.h1("Sources", new_page=True)
    d.para("All facts in section 2 were read on the dates shown. Nothing here comes from private or non-public information.",
           size=9.5, color=GREY)
    dates = {}
    for f in shown:
        dates.setdefault(f["source"], f["retrieved"])
    d.table(["#", "Source", "Retrieved"], [(str(i), u, dates.get(u, "")) for i, u in enumerate(urls, 1)],
            [1.2, 13.0, 2.8], "Sources", size=8.5, first_col_bold=False)
    path = outdir / f"{rec['id'].replace('-', '_', 1)}.docx"
    d.save(path, f"FluxMuse proposal for {n}", "Tailored prospect proposal")
    return path


def build_briefs(recs, path):
    d = ProspectDoc("INTERNAL: prospect briefs. Do not send.", "internal")
    d.footer_text = "Fluxmuse Pty Ltd · INTERNAL prospect briefs · do not send to prospects"
    d.cover_page("Internal", "Prospect briefs", f"{len(recs)} prospects · {AS_AT}",
                 [("Purpose", "How to approach each prospect, what to say, and where the doubts are."),
                  ("Read with", "Outreach_Playbook.md and Prospect_List.xlsx"),
                  ("Rule", "Human-sent, one-to-one, published channels only. See CURRENT_OFFER.md section 6.")],
                 "Internal working document. It contains candid fit assessments and must not be shared with the businesses named in it.")
    for i, r in enumerate(recs):
        d.h1(f"{r['id'][:2]}. {r['name']}", new_page=(i > 0))
        t = TIERS[r["recommended_plan"]["tier"]]
        d.kv_table([("Trade / place", f"{r['trade']} · {r['location']}"),
                    ("Wave / attainability", f"Wave {r['wave']} · {r['fit']['attainability']}/5. {r['why_wave']}"),
                    ("Plan / potential", f"{t[0]}" + (f" ({money(t[1])}/month)" if t[1] else "") + f" · {r['fit']['revenue_potential']}"),
                    ("Decision-maker", r["decision_maker"])], caption="Layout: brief summary", size=9.5)
        d.h3("How to approach")
        d.kv_table([("Route", r["approach"]["route"]), ("Where it is published", r["approach"]["contact_source"]),
                    ("Why this route", r["approach"]["compliance_note"])], caption="Layout: approach", size=9.5)
        d.callout("Opening message (a person sends this)", [r["approach"]["opening_message"]], fill=TINT, accent="FF6A00")
        d.h3("Fit, honestly")
        d.para(r["fit"]["reasoning"])
        d.h3("What to validate on the call")
        d.bullets([f"{p['pain']} Ask: {p['validate_with']}" for p in r["pains"]])
        d.h3("What we could not confirm")
        d.bullets(r["gaps"] or ["Nothing outstanding."])
    d.save(path, "FluxMuse internal prospect briefs", "Internal")


def build_xlsx(recs, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ORANGE = PatternFill("solid", fgColor="FF6A00")
    INPUT = PatternFill("solid", fgColor="FFF4CC")
    hb = Font(bold=True, color="20242B")
    wb = Workbook()

    ws = wb.active
    ws.title = "Prospects"
    cols = ["Rank", "ID", "Prospect", "Trade", "Segment", "Location", "Wave", "Attainability (1-5)", "Plan", "List price / month (R)",
            "Potential", "Decision-maker", "Approach route", "Where the route is published", "Top hypothesis to test",
            "Proposal file", "Status", "Next action", "Next action date", "Notes"]
    ws.append(cols)
    ordered = sorted(recs, key=lambda r: (r["wave"], -int(r["fit"]["attainability"]), -(TIERS[r["recommended_plan"]["tier"]][1] or 0)))
    for rank, r in enumerate(ordered, 1):
        t = TIERS[r["recommended_plan"]["tier"]]
        ws.append([rank, r["id"], r["name"], r["trade"], r["segment"], r["location"], r["wave"], int(r["fit"]["attainability"]),
                   t[0], t[1] or 0, r["fit"]["revenue_potential"], r["decision_maker"], r["approach"]["route"],
                   r["approach"]["contact_source"], r["pains"][0]["pain"], f"proposals/{r['id'].replace('-', '_', 1)}.docx",
                   "To contact", "", None, ""])
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = ORANGE, hb
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate([6, 16, 26, 26, 10, 22, 6, 10, 10, 12, 26, 32, 40, 38, 44, 34, 14, 30, 14, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1='"To contact,Contacted,Replied,Call booked,Proposal sent,Won,Lost,Parked"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"Q2:Q{len(recs) + 1}")
    for r_i in range(2, len(recs) + 2):
        cell = ws.cell(row=r_i, column=16)
        cell.hyperlink = cell.value
        cell.font = Font(color="C24E00", underline="single")
        ws.cell(row=r_i, column=19).number_format = "yyyy-mm-dd"

    # pipeline maths: inputs in yellow, everything else is a formula
    pm = wb.create_sheet("Pipeline maths")
    pm["A1"] = "Pipeline maths: what this list could plausibly produce"
    pm["A1"].font = Font(bold=True, size=13)
    pm["A2"] = "Yellow cells are assumptions to replace with real reply and close rates as soon as you have them. Nothing here is a forecast."
    pm.append([])
    pm.append(["Wave", "Prospects", "Reply rate", "Call held (of replies)", "Buy (of calls)", "Expected customers", "Avg list price / month (R)", "Expected MRR (R)"])
    for c in range(1, 9):
        pm.cell(row=4, column=c).fill, pm.cell(row=4, column=c).font = ORANGE, hb
    assumptions = {1: (0.25, 0.6, 0.3), 2: (0.2, 0.5, 0.25), 3: (0.15, 0.5, 0.2)}
    for w in (1, 2, 3):
        row = 4 + w
        pm.cell(row=row, column=1, value=f"Wave {w}")
        pm.cell(row=row, column=2, value=f"=COUNTIF(Prospects!G:G,{w})")
        for j, v in enumerate(assumptions[w]):
            c = pm.cell(row=row, column=3 + j, value=v)
            c.fill, c.number_format = INPUT, "0%"
        pm.cell(row=row, column=6, value=f"=B{row}*C{row}*D{row}*E{row}").number_format = "0.0"
        pm.cell(row=row, column=7, value=f"=IFERROR(AVERAGEIF(Prospects!G:G,{w},Prospects!J:J),0)").number_format = "#,##0"
        pm.cell(row=row, column=8, value=f"=F{row}*G{row}").number_format = "#,##0"
    pm.cell(row=8, column=1, value="Total").font = hb
    pm.cell(row=8, column=2, value="=SUM(B5:B7)")
    pm.cell(row=8, column=6, value="=SUM(F5:F7)").number_format = "0.0"
    pm.cell(row=8, column=8, value="=SUM(H5:H7)").number_format = "#,##0"
    pm["A10"] = ("Read this honestly: the 'Expected customers' total in F8 is the headline. At typical cold, pay-first rates it is "
                 "around one paying customer from this whole list, not twenty. Demos on the prospect's own catalogue and warm "
                 "introductions convert better; replace the yellow cells with your real rates as soon as you have them.")
    pm["A10"].alignment = Alignment(wrap_text=True, vertical="top")
    pm.merge_cells("A10:H10")
    pm.row_dimensions[10].height = 58
    pm["A12"] = "End-to-end conversion per prospect"
    pm["F12"] = "=IFERROR(F8/B8,0)"
    pm["F12"].number_format = "0.0%"
    pm["A13"] = "Prospects to approach for 3 paying customers at these rates"
    pm["F13"] = '=IFERROR(3/F12,"n/a")'
    pm["F13"].number_format = "0"
    for c in ("A12", "A13"):
        pm[c].font = hb
    for i, w in enumerate([12, 11, 11, 20, 14, 18, 24, 18], 1):
        pm.column_dimensions[get_column_letter(i)].width = w

    src = wb.create_sheet("Sources")
    src.append(["Prospect", "Fact", "Source", "Retrieved"])
    for c in range(1, 5):
        src.cell(row=1, column=c).fill, src.cell(row=1, column=c).font = ORANGE, hb
    for r in ordered:
        for f in r["public_facts"]:
            src.append([r["name"], f["fact"], f["source"], f["retrieved"]])
    for i, w in enumerate([28, 90, 60, 12], 1):
        src.column_dimensions[get_column_letter(i)].width = w
    for row in src.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    src.freeze_panes = "A2"

    rules = wb.create_sheet("Outreach rules")
    for line in ["Human-sent, one-to-one messages only. No bulk sends, no automated DMs, no AI cold calls (POPIA s69).",
                 "Use only a route the business publishes itself (website form, listed business email or number, public business account) or a warm introduction.",
                 "Never guess or scrape an email address or a personal number. Public figures: management or booking.",
                 "Every first message: who we are, why this business (one true observation), the ask (a 20-minute call), an opt-out line.",
                 "No trial, free-period, pilot, customer-count or results language. FluxMuse has no customers or results yet, say so.",
                 "If someone says stop, stop, and record it in the Status column as Lost with a note."]:
        rules.append([line])
    rules.column_dimensions["A"].width = 150
    for row in rules.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True)
    wb.save(path)


# ---- entry points ----------------------------------------------------------------------------------------
def load():
    recs, bad = [], False
    for f in sorted(DATA.glob("[0-9][0-9]-*.json")):
        try:
            rec = json.loads(f.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            print(f"FAIL {f.name}: invalid JSON: {e}")
            bad = True
            continue
        E, W = validate(rec, f.name)
        status = "FAIL" if E else ("warn" if W else "ok  ")
        print(f"{status} {f.name}")
        for e in E:
            print(f"       error: {e}")
        for w in W:
            print(f"       warn : {w}")
        bad |= bool(E)
        recs.append(rec)
    return recs, bad


SELFTEST = {
    "id": "99-example-studio", "slot": "test", "name": "Example Studio", "legal_or_trading_name": "", "segment": "solo",
    "trade": "Hair braiding studio", "location": "Soweto, Johannesburg", "wave": 1, "why_wave": "Owner decides alone.",
    "one_line_pitch": "Turn the photos you already take into a shop link and a WhatsApp order flow, set up with you by hand.",
    "public_facts": [{"fact": f"Synthetic fact number {i}.", "source": f"https://example.com/{i}", "retrieved": "2026-09-26"} for i in range(1, 7)],
    "businesses": [], "channels_observed": ["Instagram business account (public)"],
    "decision_maker": "Owner.",
    "approach": {"route": "Public business Instagram account, sent by a person.", "contact_source": "https://example.com/1",
                 "opening_message": "Hi, I'm from FluxMuse, a new WhatsApp tool for small businesses. I noticed your prices are on Instagram highlights and bookings come by DM. "
                                    "We help turn photos into a simple shop link and booking chat. Could we have a 20-minute call this week? "
                                    "If you'd rather I didn't message again, just say so and I'll stop.",
                 "compliance_note": "One-to-one message via the business's own public channel."},
    "pains": [{"pain": f"Pain {i}.", "basis": "inferred", "evidence": "Typical for the trade.", "validate_with": f"Question {i}?"} for i in range(1, 4)],
    "fit": {"attainability": 4, "revenue_potential": "R289/month (Micro)", "reasoning": "Good fit. Could be poor if they have no WhatsApp."},
    "recommended_plan": {"tier": "micro", "why": "One studio, one number.", "upgrade_path": "Starter when a second stylist joins."},
    "what_we_would_do": [{"capability": f"Capability {i}", "status": "live" if i % 2 else "in_setup", "for_them": "It helps."} for i in range(1, 5)],
    "first_30_days": [{"week": w, "actions": [f"Action for week {w}"]} for w in (1, 2, 3, 4)],
    "we_would_track": ["Conversations from the shop link", "Bookings by WhatsApp", "Time spent on replies"],
    "objections": [{"objection": f"Objection {i}?", "response": "Answer."} for i in range(1, 4)],
    "related_reading": [], "gaps": ["None."],
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        E, W = validate(SELFTEST)
        assert not E, E
        tmp = Path(tempfile.mkdtemp())
        print("proposal:", build_proposal(SELFTEST, tmp))
        build_briefs([SELFTEST], tmp / "briefs.docx")
        build_xlsx([SELFTEST], tmp / "list.xlsx")
        print("selftest outputs in", tmp, "warnings:", W)
        return 0
    recs, bad = load()
    if a.check or bad:
        return 1 if bad else 0
    (OUT / "proposals").mkdir(parents=True, exist_ok=True)
    for r in recs:
        print("built", build_proposal(r, OUT / "proposals").name)
    build_briefs(recs, OUT / "Prospect_Briefs.docx")
    build_xlsx(recs, OUT / "Prospect_List.xlsx")
    print(f"built {len(recs)} proposals + briefs + list")
    return 0


if __name__ == "__main__":
    sys.exit(main())
